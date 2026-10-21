"""
Toko X (= channel BLOOMIE) — Hitung ulang berbasis SKU INDUK (Brief 2).

Melanjutkan tokox_analisis_lanjutan.py. Perbedaan utama vs skrip lama:
  1. Grosir/eceran diklasifikasi dari qty per SKU INDUK dalam satu nota
     (bukan total qty per nota). 12 pcs satu produk = grosir; 12 produk @1 = eceran.
  2. Item_No yang hanya beda kode warna digabung jadi satu SKU induk
     (akhiran di belakang Item_No), divalidasi lewat deskripsi.

Data mentah dibaca dari SQLite lokal (channel BLOOMIE) via helper lama.
Output: ~/Bloomie_Analysis/TokoX_Analisis_SKU.xlsx

  python analysis/tokox_analisis_sku.py [--reopen 2026-07-19] [--out DIR]

Catatan: "Toko X" = BLOOMIE (kerahasiaan). Aturan CLAUDE.md tetap berlaku KECUALI
definisi SKU, yang untuk analisis Toko X diganti jadi SKU induk (lihat 01_SKU_Mapping).
"""
import os, re, argparse
from collections import defaultdict, Counter
from datetime import datetime, date, timedelta
from statistics import median
import openpyxl
import _bloomie_common as B

BULK = {"BALL","BOX","CTN","DUS","GROS","GROSS","KARTON","KODI","LSN","LUSIN","PACK","PAK","RIM"}
GENERIK_NAMES = {"UMUM","CASH","WALK IN","WALKIN","GUEST","-","","TUNAI","UMUM/CASH"}
WALKIN = {"JKT-00001", "JKT-00002"}          # akun kasir toko & "MEMBER" (temuan brief 1)
OUTLIER_AKUN = "JKT-00238"                    # nota outlier 31 Mar 2026 (42.524 pcs)

# kata warna (ID + EN) — dibuang saat membandingkan deskripsi varian
COLORS = {
    "BLACK","WHITE","CREAM","KREM","PINK","RED","BLUE","GREEN","YELLOW","PURPLE","BROWN",
    "GREY","GRAY","SILVER","GOLD","NAVY","MAROON","BEIGE","ORANGE","KHAKI","TOSCA","TOSKA",
    "LILAC","LAVENDER","MINT","PEACH","CORAL","MUSTARD","OLIVE","TURQUOISE","MAGENTA","VIOLET",
    "IVORY","BURGUNDY","FUCHSIA","FUSCHIA","TEAL","CYAN","NUDE","MOCHA","COFFEE","CHOCO",
    "CHOCOLATE","CHOCOLAT","ROSE","SALMON","APRICOT","CHAMPAGNE","MILO","TAUPE","DENIM","CARAMEL",
    "HITAM","PUTIH","MERAH","BIRU","HIJAU","KUNING","UNGU","COKLAT","COKELAT","ABU","EMAS","PERAK",
    "DONGKER","TOSCADONGKER","MUDA","TUA","DARK","LIGHT","LT","DK","MOTIF","POLOS","BENING","TRANSPARAN",
}
# prefix kode akun -> wilayah (dari pengamatan; default = prefix apa adanya)
WILAYAH = {
    "JKT":"DKI Jakarta","BKS":"Bekasi","TNG":"Tangerang","BGR":"Bogor","DPK":"Depok","BDG":"Bandung",
    "SBY":"Surabaya","SMG":"Semarang","SOLO":"Solo","YGY":"Yogyakarta","JOG":"Yogyakarta","MLG":"Malang",
    "MDN":"Medan","PLB":"Palembang","PKU":"Pekanbaru","PDG":"Padang","LPG":"Lampung","DPS":"Bali",
    "MKS":"Makassar","MKR":"Makassar","BPN":"Balikpapan","BJM":"Banjarmasin","PNK":"Pontianak","SMD":"Samarinda",
}


def qf(v):
    try: return float(v)
    except Exception: return 0.0


def tier_of(q):
    if q <= 1: return "T1"
    if q <= 5: return "T2"
    if q <= 11: return "T3a"
    if q <= 20: return "T3b"
    if q <= 50: return "T4"
    return "T5"
TIER_ORDER = ["T1","T2","T3a","T3b","T4","T5"]

SEG_ORDER = ["S1","S2","S3","S4","S5"]
SEG_LABEL = {"S1":"Eceran","S2":"Eceran borong","S3":"Setengah lusin",
             "S4":"Grosir (1-2 lusin/SKU)","S5":"Grosir besar (>=3 lusin/SKU)"}


def seg_of(max_qty, bulk):
    """Segmen nota dari qty terbesar satu SKU induk. Satuan bulk -> minimal S4."""
    if max_qty >= 36: return "S5"
    if max_qty >= 12 or bulk: return "S4"
    if max_qty >= 6: return "S3"
    if max_qty >= 3: return "S2"
    return "S1"


def month_bucket(d, reopen):
    if d.year == reopen.year and d.month == reopen.month:
        return f"{d.year}-{d.month:02d}a" if d.day < reopen.day else f"{d.year}-{d.month:02d}b"
    return f"{d.year}-{d.month:02d}"


def mkey(m):
    y = int(m[:4]); mo = int(m[5:7]); half = 0 if m.endswith("a") else (1 if m.endswith("b") else -1)
    return (y, mo, half)


def per30(total, days):
    return total / days * 30 if days else 0


def wilayah_of(cust):
    pre = (cust or "").split("-")[0].strip().upper()
    return WILAYAH.get(pre, pre or "—")


# ---------------- SKU induk ----------------
def reduce_item(item):
    """(induk_kandidat, pola). pola: A huruf1-3, B angka2digit, C huruf-nempel, D standar."""
    item = (item or "").strip()
    segs = item.split("-")
    last = segs[-1] if segs else ""
    if len(segs) >= 4 and re.fullmatch(r"[A-Za-z]{1,3}", last):
        return "-".join(segs[:-1]), "A"
    if len(segs) == 5 and re.fullmatch(r"\d{2}", last):
        return "-".join(segs[:-1]), "B"
    if re.fullmatch(r"\d+[A-Za-z]+", last):
        return item, "C"
    return item, "D"


def norm_desc(s):
    """Uppercase, buang tanda baca & token warna -> untuk banding varian."""
    s = (s or "").upper()
    toks = re.split(r"[^A-Z0-9]+", s)
    toks = [t for t in toks if t and t not in COLORS]
    return " ".join(toks)


def build_sku_mapping(items_desc):
    """items_desc: {item_no: deskripsi}. Return {item_no: dict(sku_induk,pola,digabung,alasan,perlu_review)}."""
    groups = defaultdict(list)                       # induk_kandidat -> [item_no]
    pola_of = {}
    for it in items_desc:
        induk_c, pola = reduce_item(it)
        pola_of[it] = (induk_c, pola)
        groups[induk_c].append(it)

    POLA_LABEL = {"A":"huruf_1-3","B":"angka_2digit","C":"huruf_nempel","D":"standar_4segmen"}
    mapping = {}
    for induk_c, members in groups.items():
        variants = [m for m in members if pola_of[m][0] != m]        # anggota yang punya akhiran (A/B)
        has_variant = len(variants) > 0
        variant_polas = {pola_of[m][1] for m in variants}
        stripped = {norm_desc(items_desc.get(m, "")) for m in members}

        if not has_variant:
            merge = False; review = False; alasan = "SKU induk (tak ada varian akhiran)"
        elif len(stripped) == 1:
            # deskripsi cocok setelah warna dibuang
            lone_b = (variant_polas == {"B"}) and (len(members) == 1)
            if lone_b:
                # akhiran -NN tunggal tanpa sibling: indeks item master, bukan kode warna.
                # Tidak digabung & bukan kasus ambigu -> SKU standalone.
                merge = False; review = False
                alasan = "akhiran angka 2 digit varian tunggal (tanpa sibling) -> SKU standalone, tidak digabung"
            else:
                merge = True; review = False
                alasan = "deskripsi identik/hanya beda kata warna -> digabung"
        else:
            merge = False; review = True
            alasan = "deskripsi beda di luar warna (cek ukuran/model/isi) -> jangan digabung"

        for m in members:
            induk_m, pola_m = pola_of[m]
            if pola_m in ("A", "B") and merge:
                mapping[m] = {"sku_induk": induk_c, "pola": POLA_LABEL[pola_m],
                              "digabung": (induk_c != m), "alasan": alasan, "perlu_review": False}
            elif pola_m in ("A", "B"):
                mapping[m] = {"sku_induk": m, "pola": POLA_LABEL[pola_m],
                              "digabung": False, "alasan": alasan, "perlu_review": True}
            elif pola_m == "C":
                mapping[m] = {"sku_induk": m, "pola": POLA_LABEL["C"],
                              "digabung": False, "alasan": "huruf menempel tanpa '-' -> tidak dipisah, hanya dilaporkan",
                              "perlu_review": False}
            else:
                mapping[m] = {"sku_induk": m, "pola": POLA_LABEL["D"],
                              "digabung": False, "alasan": "format 4 segmen standar", "perlu_review": False}
    return mapping


# ================= Perbaikan Brief 3 (hanya dipakai di --v2) =================
CONV_PCS = {"LSN": 12, "LUSIN": 12, "KODI": 20, "GROS": 144, "GROSS": 144}   # dikonversi ke pcs
KEMASAN = {"BOX","PAK","PACK","DUS","CTN","KARTON","BALL","RIM","KG"}        # TIDAK auto-grosir
_SIZE_RE = re.compile(r"(UK\s*\d+|\d+\s*CM|\d+\s*MM|\(\s*\d+\s*(?:PCS|PC|PIECES|SET)\s*\)|\d+\s*(?:BOX|SET|PAK|PACK|LUSIN|LSN))")


def seg_v2(mx):
    """Segmen tanpa paksaan satuan bulk (Perbaikan 1). Qty lusinan sudah dikonversi ke pcs."""
    if mx >= 36: return "S5"
    if mx >= 12: return "S4"
    if mx >= 6: return "S3"
    if mx >= 3: return "S2"
    return "S1"


def size_content_sig(desc):
    """Tanda tangan ukuran/isi/kemasan dari deskripsi -> pengecualian GABUNG_SEMUA."""
    return frozenset(re.sub(r"\s+", "", m).upper() for m in _SIZE_RE.findall((desc or "").upper()))


def build_gabung_mapping(items_desc):
    """Skenario GABUNG_SEMUA: semua item_no dgn induk 4 segmen sama digabung,
    KECUALI deskripsi beda ukuran/isi (sub-kelompok per tanda tangan ukuran)."""
    groups = defaultdict(list)
    for it in items_desc:
        segs = it.split("-")
        induk4 = "-".join(segs[:4]) if len(segs) >= 4 else it
        groups[induk4].append(it)
    mp = {}
    for induk4, members in groups.items():
        for m in members:
            sig = size_content_sig(items_desc.get(m, ""))
            mp[m] = induk4 if not sig else induk4 + "|" + "+".join(sorted(sig))
    return mp


def load_clean(reopen):
    """Baca + bersihkan data mentah (identik Brief 1/2). Return None bila kosong."""
    raw = B.fetch_bloomie()
    print(f"{len(raw)} baris mentah Toko X.")
    if not raw:
        return None
    dropped_retur = dropped_qty0 = dropped_nodate = unit_empty = 0
    rows = []; desc_count = defaultdict(Counter)
    for r in raw:
        d = B.d_(r.get("posting_date"))
        if not d:
            dropped_nodate += 1; continue
        q = qf(r.get("quantity")); rev = qf(r.get(B.REV))
        if rev < 0:
            dropped_retur += 1; continue
        if q <= 0:
            dropped_qty0 += 1; continue
        u = (r.get("unit") or "").strip().upper()
        if u == "PIECES": u = "PCS"
        if u == "": unit_empty += 1; u = "—"
        item = (r.get("item_no") or "").strip()
        dsc = (r.get("item_description") or "").strip()
        if dsc:
            desc_count[item][dsc] += 1
        rows.append({"doc": r.get("document_number"), "date": d, "cust": (r.get("customer_code") or "").strip(),
                     "name": r.get("customer_name"), "item": item, "kat": (r.get("kategori") or "Lainnya"),
                     "qty": q, "unit": u, "rev": rev})
    last_date = max(r["date"] for r in rows)
    DAYS = {"PRA": 199, "PASCA": (last_date - reopen).days + 1, "PRA_BASELINE": 181}
    jun30 = date(reopen.year, 6, 30)
    for r in rows:
        r["periode"] = "PRA" if r["date"] < reopen else "PASCA"
        r["bulan"] = month_bucket(r["date"], reopen)
        r["baseline"] = (r["date"] <= jun30)
    items_desc = {it: c.most_common(1)[0][0] if c else "" for it, c in desc_count.items()}
    for r in rows:
        items_desc.setdefault(r["item"], "")
    stats = {"dropped_retur": dropped_retur, "dropped_qty0": dropped_qty0,
             "dropped_nodate": dropped_nodate, "unit_empty": unit_empty}
    return rows, items_desc, desc_count, last_date, DAYS, jun30, stats


def build_v2(rows, items_desc, last_date, DAYS, jun30, reopen, outdir):
    """Brief 3: Perbaikan 1 (konversi lusinan, kemasan bukan auto-grosir) + Perbaikan 2
    (skenario GABUNG_SEMUA) + sheet C1-C4. Output TokoX_Analisis_SKU_v2.xlsx."""
    jun1 = date(reopen.year, 6, 1)
    sku_def = {it: build_sku_mapping(items_desc)[it]["sku_induk"] for it in items_desc}
    sku_gab = build_gabung_mapping(items_desc)

    def eff(unit, qty):                      # Perbaikan 1: lusinan -> pcs
        if unit in CONV_PCS: return "PCS", qty * CONV_PCS[unit]
        return unit, qty

    def seg_calc(mx, units, bulk_force):
        if bulk_force and (units & BULK):    # replikasi aturan Brief 2 (bulk -> minimal S4)
            base = seg_v2(mx)
            return "S4" if base in ("S1", "S2", "S3") else base
        return seg_v2(mx)

    def build(sku_of, perbaikan1):
        notas = {}
        for r in rows:
            n = notas.get(r["doc"])
            if n is None:
                n = notas[r["doc"]] = {"date": r["date"], "cust": r["cust"], "name": r["name"],
                                       "rev": 0.0, "qty": 0.0, "units": set(), "skus": set(),
                                       "line": defaultdict(float), "catrev": defaultdict(float), "kemasan": False}
            elif r["date"] < n["date"]:
                n["date"] = r["date"]
            sku = sku_of[r["item"]]
            u, q = eff(r["unit"], r["qty"]) if perbaikan1 else (r["unit"], r["qty"])
            n["rev"] += r["rev"]; n["qty"] += r["qty"]
            n["units"].add(r["unit"]); n["skus"].add(sku)
            n["line"][(sku, u)] += q
            n["catrev"][r["kat"]] += r["rev"]
            if r["unit"] in KEMASAN: n["kemasan"] = True
        for n in notas.values():
            n["periode"] = "PRA" if n["date"] < reopen else "PASCA"
            mx = max(n["line"].values()) if n["line"] else 0
            n["max_qty_sku"] = mx; n["n_sku"] = len(n["skus"])
            n["seg"] = seg_calc(mx, n["units"], bulk_force=not perbaikan1)
            n["grosir"] = n["seg"] in ("S4", "S5")
            n["walkin"] = n["cust"] in WALKIN
        return notas

    def grp_of(n): return "WALKIN" if n["walkin"] else "AKUN"

    def cust_tipe(notas):
        cp = defaultdict(lambda: {"n": 0, "g": 0})
        for n in notas.values():
            if n["cust"]:
                e = cp[(n["cust"], n["periode"])]; e["n"] += 1
                if n["grosir"]: e["g"] += 1
        def tipe(c, p):
            if c in WALKIN: return "Walk-in"
            e = cp.get((c, p))
            if not e or not e["n"]: return ""
            r = e["g"] / e["n"]
            return "Grosir" if r >= 0.8 else ("Campuran" if r >= 0.2 else "Eceran")
        return cp, tipe

    def grosir_summary(notas):
        out = {}
        for p in ("PRA", "PASCA"):
            for grp in ("AKUN", "WALKIN"):
                pn = [n for n in notas.values() if n["periode"] == p and grp_of(n) == grp]
                totn = len(pn) or 1; totr = sum(n["rev"] for n in pn) or 1
                gn = [n for n in pn if n["grosir"]]
                out[(p, grp)] = {"nota": len(pn), "pct_nota_grosir": len(gn)/totn,
                                 "pct_rev_grosir": sum(n["rev"] for n in gn)/totr}
        return out

    def tipe_dist(notas):
        cp, tipe = cust_tipe(notas)
        out = defaultdict(lambda: defaultdict(int))     # periode -> tipe -> jumlah akun
        for p in ("PRA", "PASCA"):
            custs = {c for (c, pp) in cp if pp == p}
            for c in custs:
                out[p][tipe(c, p)] += 1
        return out

    def winback(notas):
        cp, tipe = cust_tipe(notas)
        first = {}; last = {}; name = {}; cat = defaultdict(lambda: defaultdict(float))
        rev_pra = defaultdict(float); nota_pra = defaultdict(int); sku_pra = defaultdict(set)
        per = defaultdict(set)
        for n in notas.values():
            c = n["cust"]
            if not c: continue
            first[c] = n["date"] if c not in first else min(first[c], n["date"])
            last[c] = n["date"] if c not in last else max(last[c], n["date"])
            name.setdefault(c, n["name"]); per[c].add(n["periode"])
            if n["periode"] == "PRA":
                rev_pra[c] += n["rev"]; nota_pra[c] += 1; sku_pra[c] |= n["skus"]
                for k, v in n["catrev"].items(): cat[c][k] += v
        out = []
        for c in rev_pra:
            if "PASCA" in per[c] or c in WALKIN: continue
            if last[c] < jun1: continue
            tp = tipe(c, "PRA")
            if tp not in ("Grosir", "Campuran"): continue
            topcat = max(cat[c].items(), key=lambda x: x[1])[0] if cat[c] else ""
            out.append({"kode_customer": c, "nama_customer": name.get(c, ""), "wilayah": wilayah_of(c),
                        "tipe_pembeli_pra": tp, "nota_pra": nota_pra[c], "revenue_pra": round(rev_pra[c]),
                        "tgl_terakhir_beli": last[c].isoformat(), "kategori_utama_pra": topcat,
                        "jumlah_sku_induk_unik_pra": len(sku_pra[c])})
        out.sort(key=lambda x: -x["revenue_pra"])
        return out

    # ---- bangun semua skenario ----
    nd_old = build(sku_def, False)       # DEFAULT tanpa Perbaikan 1 (= Brief 2) -> validasi 143
    nd = build(sku_def, True)            # DEFAULT + Perbaikan 1 -> dasar sheet C
    gab = build(sku_gab, True)           # GABUNG_SEMUA + Perbaikan 1
    _, tipe_nd = cust_tipe(nd)
    gs_nd = grosir_summary(nd); gs_gab = grosir_summary(gab)
    wb_def = winback(nd); wb_old = winback(nd_old); wb_gab = winback(gab)

    # sku first-sold (DEFAULT) utk status baru/lama
    sku_first = {}
    for r in rows:
        s = sku_def[r["item"]]
        sku_first[s] = r["date"] if s not in sku_first else min(sku_first[s], r["date"])

    # ================= WORKBOOK =================
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    def sheet(name, cols, data):
        ws = wb.create_sheet(name[:31]); ws.append(cols)
        for row in data: ws.append([row.get(c, "") for c in cols])
        return ws

    # ---------- 00_README ----------
    rd = []
    def rm(i, v, c=""): rd.append({"item": i, "nilai": v, "catatan": c})
    rm("toko", "Toko X", "skenario v2 (Brief 3)")
    for p in ("PRA", "PASCA"):
        sub = [r for r in rows if r["periode"] == p]
        rm(f"revenue_{p}", round(sum(r["rev"] for r in sub)))
        rm(f"nota_{p}", len({r["doc"] for r in sub}))
        rm(f"customer_unik_{p}", len({r["cust"] for r in sub if r["cust"]}))
    rm("target_brief1_2", "PRA 9.339.437.720/10.860 ; PASCA 6.622.181.617/9.767", "harus tetap sama")
    # ringkasan wajib Brief 3
    rm("P1_pct_nota_grosir_walkin_PRA_sebelum", round(grosir_summary(nd_old)[("PRA","WALKIN")]["pct_nota_grosir"], 4))
    rm("P1_pct_nota_grosir_walkin_PRA_sesudah", round(gs_nd[("PRA","WALKIN")]["pct_nota_grosir"], 4))
    rm("P1_pct_nota_grosir_walkin_PASCA_sebelum", round(grosir_summary(nd_old)[("PASCA","WALKIN")]["pct_nota_grosir"], 4))
    rm("P1_pct_nota_grosir_walkin_PASCA_sesudah", round(gs_nd[("PASCA","WALKIN")]["pct_nota_grosir"], 4))
    for p in ("PRA", "PASCA"):
        tot_old = sum(n["rev"] for n in nd_old.values() if n["periode"] == p) or 1
        g_old = sum(n["rev"] for n in nd_old.values() if n["periode"] == p and n["grosir"])
        tot_new = sum(n["rev"] for n in nd.values() if n["periode"] == p) or 1
        g_new = sum(n["rev"] for n in nd.values() if n["periode"] == p and n["grosir"])
        rm(f"P1_pct_revenue_grosir_{p}_sebelum", round(g_old/tot_old, 4))
        rm(f"P1_pct_revenue_grosir_{p}_sesudah", round(g_new/tot_new, 4))
    rm("sku_induk_DEFAULT", len(set(sku_def.values())))
    rm("sku_induk_GABUNG_SEMUA", len(set(sku_gab.values())))
    rm("winback_DEFAULT_tanpa_P1", f"{len(wb_old)} akun / Rp{round(sum(w['revenue_pra'] for w in wb_old)):,}",
       "target validasi brief: 143 akun / Rp1.182.581.000")
    rm("winback_DEFAULT_dgn_P1", f"{len(wb_def)} akun / Rp{round(sum(w['revenue_pra'] for w in wb_def)):,}")
    rm("winback_GABUNG_SEMUA_dgn_P1", f"{len(wb_gab)} akun / Rp{round(sum(w['revenue_pra'] for w in wb_gab)):,}")
    sku_baru = {s for s, d in sku_first.items() if d >= reopen}
    for grp in ("AKUN", "WALKIN"):
        base = sum(r["rev"] for r in rows if r["periode"] == "PASCA" and grp_of(nd[r["doc"]]) == grp) or 1
        rb = sum(r["rev"] for r in rows if r["periode"] == "PASCA" and grp_of(nd[r["doc"]]) == grp and sku_def[r["item"]] in sku_baru)
        rm(f"share_revenue_sku_baru_{grp}", round(rb/base, 4))
    sheet("00_README", ["item", "nilai", "catatan"], rd)

    # ---------- A_Perbaikan1_Shift ----------
    shift = defaultdict(lambda: {"pindah": 0, "total": 0})
    for doc, n in nd.items():
        key = (n["periode"], grp_of(n)); shift[key]["total"] += 1
        if nd_old[doc]["grosir"] != n["grosir"]: shift[key]["pindah"] += 1
    dA = []
    for (p, grp) in sorted(shift):
        e = shift[(p, grp)]; g = gs_nd[(p, grp)]
        dA.append({"periode": p, "kelompok": grp, "nota_pindah_grosir": e["pindah"], "total_nota": e["total"],
                   "pct_nota_pindah": round(e["pindah"]/e["total"], 4) if e["total"] else "",
                   "pct_nota_grosir_setelah_P1": round(g["pct_nota_grosir"], 4),
                   "pct_revenue_grosir_setelah_P1": round(g["pct_rev_grosir"], 4)})
    sheet("A_Perbaikan1_Shift", ["periode","kelompok","nota_pindah_grosir","total_nota","pct_nota_pindah",
                                 "pct_nota_grosir_setelah_P1","pct_revenue_grosir_setelah_P1"], dA)

    # ---------- B_Skenario (DEFAULT vs GABUNG_SEMUA) ----------
    dB = [{"blok": "sku_induk", "dimensi": "jumlah", "DEFAULT": len(set(sku_def.values())),
           "GABUNG_SEMUA": len(set(sku_gab.values()))}]
    for (p, grp) in sorted(gs_nd):
        dB.append({"blok": "pct_nota_grosir", "dimensi": f"{p}/{grp}",
                   "DEFAULT": round(gs_nd[(p, grp)]["pct_nota_grosir"], 4),
                   "GABUNG_SEMUA": round(gs_gab[(p, grp)]["pct_nota_grosir"], 4)})
        dB.append({"blok": "pct_revenue_grosir", "dimensi": f"{p}/{grp}",
                   "DEFAULT": round(gs_nd[(p, grp)]["pct_rev_grosir"], 4),
                   "GABUNG_SEMUA": round(gs_gab[(p, grp)]["pct_rev_grosir"], 4)})
    td_nd = tipe_dist(nd); td_gab = tipe_dist(gab)
    for p in ("PRA", "PASCA"):
        for t in ("Grosir", "Campuran", "Eceran", "Walk-in"):
            dB.append({"blok": "tipe_akun", "dimensi": f"{p}/{t}",
                       "DEFAULT": td_nd[p].get(t, 0), "GABUNG_SEMUA": td_gab[p].get(t, 0)})
    dB.append({"blok": "winback", "dimensi": "jumlah_akun", "DEFAULT": len(wb_def), "GABUNG_SEMUA": len(wb_gab)})
    sheet("B_Skenario", ["blok", "dimensi", "DEFAULT", "GABUNG_SEMUA"], dB)

    # ---------- C1_SKU_Baru ----------
    c1 = defaultdict(lambda: {"rev": 0.0, "qty": 0.0, "docs": set(), "custs": set()})
    for r in rows:
        if r["periode"] != "PASCA": continue
        n = nd[r["doc"]]; s = sku_def[r["item"]]
        key = ("baru" if s in sku_baru else "lama", grp_of(n), tipe_nd(r["cust"], "PASCA") or "-", n["seg"])
        e = c1[key]; e["rev"] += r["rev"]; e["qty"] += r["qty"]; e["docs"].add(r["doc"])
        if r["cust"]: e["custs"].add(r["cust"])
    dC1 = [{"sku_status": k[0], "kelompok": k[1], "tipe_akun": k[2], "segmen": k[3],
            "revenue": round(v["rev"]), "qty": round(v["qty"]), "jumlah_nota": len(v["docs"]),
            "jumlah_akun": len(v["custs"])} for k, v in sorted(c1.items(), key=lambda x: -x[1]["rev"])]
    sheet("C1_SKU_Baru", ["sku_status","kelompok","tipe_akun","segmen","revenue","qty","jumlah_nota","jumlah_akun"], dC1)

    # ---------- C2_Kategori_x_Kelompok ----------
    c2 = defaultdict(lambda: {"rev": 0.0, "rev_gros": 0.0})
    for r in rows:
        n = nd[r["doc"]]; key = (r["periode"], r["kat"], grp_of(n))
        e = c2[key]; e["rev"] += r["rev"]
        if n["grosir"]: e["rev_gros"] += r["rev"]
    dC2 = []
    for (p, kat, grp), e in sorted(c2.items(), key=lambda x: (x[0][0], -x[1]["rev"])):
        dC2.append({"periode": p, "kategori": kat, "kelompok": grp, "revenue": round(e["rev"]),
                    "revenue_per30": round(per30(e["rev"], DAYS[p])),
                    "pct_revenue_grosir": round(e["rev_gros"]/e["rev"], 4) if e["rev"] else ""})
    sheet("C2_Kategori_x_Kelompok", ["periode","kategori","kelompok","revenue","revenue_per30","pct_revenue_grosir"], dC2)

    # ---------- C3_Tipe_Min2Nota (AKUN non-walkin, >=2 nota) ----------
    cp_nd, _ = cust_tipe(nd)
    pra_set = {c for (c, p) in cp_nd if p == "PRA" and c not in WALKIN}
    pas_set = {c for (c, p) in cp_nd if p == "PASCA" and c not in WALKIN}
    rev_cp = defaultdict(float)
    for n in nd.values():
        if n["cust"] and n["cust"] not in WALKIN: rev_cp[(n["cust"], n["periode"])] += n["rev"]
    dC3 = []
    for p, members in (("PRA", pra_set), ("PASCA", pas_set)):
        min2 = [c for c in members if cp_nd[(c, p)]["n"] >= 2]
        by = defaultdict(list)
        for c in min2: by[tipe_nd(c, p)].append(c)
        for t, custs in sorted(by.items()):
            if p == "PRA":
                kembali = sum(1 for c in custs if c in pas_set)
                pct = round(kembali/len(custs), 4) if custs else ""
            else:
                pct = ""
            dC3.append({"periode": p, "tipe_akun": t, "jumlah_akun": len(custs),
                        "revenue": round(sum(rev_cp[(c, p)] for c in custs)), "pct_kembali": pct})
    sheet("C3_Tipe_Min2Nota", ["periode","tipe_akun","jumlah_akun","revenue","pct_kembali"], dC3)

    # ---------- C4_Winback (DEFAULT + Perbaikan 1) ----------
    sheet("C4_Winback", ["kode_customer","nama_customer","wilayah","tipe_pembeli_pra","nota_pra","revenue_pra",
                         "tgl_terakhir_beli","kategori_utama_pra","jumlah_sku_induk_unik_pra"], wb_def)

    B.autosize(wb, cap=34)
    os.makedirs(outdir, exist_ok=True)
    out = os.path.join(outdir, "TokoX_Analisis_SKU_v2.xlsx")
    wb.save(out)

    # ---- console ----
    print("\n=== VALIDASI v2 ===")
    for p in ("PRA", "PASCA"):
        sub = [r for r in rows if r["periode"] == p]
        print(f"  {p}: Rp{round(sum(r['rev'] for r in sub)):,} | nota {len({r['doc'] for r in sub})} | cust {len({r['cust'] for r in sub if r['cust']})}")
    print(f"  WINBACK DEFAULT tanpa P1: {len(wb_old)} akun / Rp{round(sum(w['revenue_pra'] for w in wb_old)):,}  (target 143 / 1.182.581.000)")
    print(f"  WINBACK DEFAULT dgn P1  : {len(wb_def)} akun / Rp{round(sum(w['revenue_pra'] for w in wb_def)):,}")
    print("\n=== RINGKASAN BRIEF 3 (tempel ke chat) ===")
    print(f"  SKU induk: DEFAULT {len(set(sku_def.values()))} vs GABUNG_SEMUA {len(set(sku_gab.values()))}")
    print(f"  %nota grosir WALKIN PRA: {grosir_summary(nd_old)[('PRA','WALKIN')]['pct_nota_grosir']:.1%} -> {gs_nd[('PRA','WALKIN')]['pct_nota_grosir']:.1%} (setelah P1)")
    for p in ("PRA", "PASCA"):
        tot_old = sum(n["rev"] for n in nd_old.values() if n["periode"] == p) or 1
        g_old = sum(n["rev"] for n in nd_old.values() if n["periode"] == p and n["grosir"])
        tot_new = sum(n["rev"] for n in nd.values() if n["periode"] == p) or 1
        g_new = sum(n["rev"] for n in nd.values() if n["periode"] == p and n["grosir"])
        print(f"  %revenue grosir {p}: {g_old/tot_old:.1%} -> {g_new/tot_new:.1%} (setelah P1)")
    for grp in ("AKUN", "WALKIN"):
        base = sum(r["rev"] for r in rows if r["periode"] == "PASCA" and grp_of(nd[r["doc"]]) == grp) or 1
        rb = sum(r["rev"] for r in rows if r["periode"] == "PASCA" and grp_of(nd[r["doc"]]) == grp and sku_def[r["item"]] in sku_baru)
        print(f"  share revenue SKU baru {grp}: {rb/base:.1%}")
    print(f"\n✓ Excel: {out}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--reopen", default="2026-07-19")
    ap.add_argument("--out", default=os.path.join(os.path.expanduser("~"), "Bloomie_Analysis"))
    ap.add_argument("--v2", action="store_true", help="Brief 3: Perbaikan 1+2 & sheet C1-C4 -> TokoX_Analisis_SKU_v2.xlsx")
    args = ap.parse_args()
    reopen = datetime.fromisoformat(args.reopen).date()

    loaded = load_clean(reopen)
    if loaded is None:
        print("Tidak ada data. Jalankan local_ingest.py dulu."); return
    rows, items_desc, desc_count, last_date, DAYS, jun30, _stats = loaded
    dropped_retur = _stats["dropped_retur"]; dropped_qty0 = _stats["dropped_qty0"]
    dropped_nodate = _stats["dropped_nodate"]; unit_empty = _stats["unit_empty"]

    if args.v2:
        build_v2(rows, items_desc, last_date, DAYS, jun30, reopen, args.out)
        return

    # ---------------- SKU INDUK MAPPING ----------------
    mapping = build_sku_mapping(items_desc)
    for r in rows:
        r["sku"] = mapping[r["item"]]["sku_induk"]

    n_item = len(items_desc)
    n_sku = len({m["sku_induk"] for m in mapping.values()})
    n_digabung = sum(1 for m in mapping.values() if m["digabung"])
    n_review = sum(1 for m in mapping.values() if m["perlu_review"])
    pola_counts = Counter(m["pola"] for m in mapping.values())

    # ---------------- NOTA-LEVEL ----------------
    notas = defaultdict(lambda: {"date": None, "cust": None, "name": None, "qty": 0.0, "rev": 0.0,
                                 "units": set(), "skus": set(), "line": defaultdict(float),
                                 "catrev": defaultdict(float)})
    for r in rows:
        n = notas[r["doc"]]
        n["date"] = r["date"] if n["date"] is None else min(n["date"], r["date"])
        if n["cust"] is None: n["cust"] = r["cust"]; n["name"] = r["name"]
        n["qty"] += r["qty"]; n["rev"] += r["rev"]
        n["units"].add(r["unit"]); n["skus"].add(r["sku"])
        n["line"][(r["sku"], r["unit"])] += r["qty"]   # qty per (sku induk, satuan) -> tak campur satuan
        n["catrev"][r["kat"]] += r["rev"]

    outlier_docs = set()
    for doc, n in notas.items():
        n["periode"] = "PRA" if n["date"] < reopen else "PASCA"
        n["bulan"] = month_bucket(n["date"], reopen)
        n["baseline"] = (n["date"] <= jun30)
        lines = n["line"]
        n["max_qty_sku"] = max(lines.values()) if lines else 0
        n["n_sku"] = len(n["skus"])
        n["bulk"] = bool(n["units"] & BULK)
        tot_q = sum(lines.values()) or 1
        n["pct_qty_lusinan"] = sum(q for q in lines.values() if q >= 12) / tot_q
        n["seg"] = seg_of(n["max_qty_sku"], n["bulk"])
        n["grosir_sku"] = n["seg"] in ("S4", "S5")
        n["tier_lama"] = tier_of(n["qty"])            # klasifikasi lama (total qty/nota)
        n["grosir_lama"] = n["bulk"] or n["qty"] >= 12
        n["walkin"] = n["cust"] in WALKIN
        if n["cust"] == OUTLIER_AKUN and n["qty"] >= 40000:
            outlier_docs.add(doc)
            n["outlier"] = True
        else:
            n["outlier"] = False

    allnotas = list(notas.values())
    kelompok_of = lambda n: "WALKIN" if n["walkin"] else "AKUN"

    # ---------------- GENERIK (laporan, dari brief 1) ----------------
    cust_nota_all = defaultdict(int); cust_rev_all = defaultdict(float)
    for n in allnotas:
        if n["cust"]: cust_nota_all[n["cust"]] += 1; cust_rev_all[n["cust"]] += n["rev"]
    generik = {}
    for c in cust_nota_all:
        nm = (c or "").strip().upper(); reason = None
        if nm in GENERIK_NAMES or c == "":
            reason = "kode/nama generik"
        else:
            nn = cust_nota_all[c]; aov = cust_rev_all[c] / nn if nn else 0
            if nn >= 50 and aov < 50000: reason = f"nota tinggi ({nn}) & AOV kecil ({round(aov)})"
        if reason: generik[c] = reason

    cust_pra = {n["cust"] for n in allnotas if n["periode"] == "PRA" and n["cust"]}
    cust_pas = {n["cust"] for n in allnotas if n["periode"] == "PASCA" and n["cust"]}

    # ================= WORKBOOK =================
    wb = openpyxl.Workbook(); wb.remove(wb.active)

    def sheet(name, columns, data):
        ws = wb.create_sheet(name[:31]); ws.append(columns)
        for row in data:
            ws.append([row.get(c, "") for c in columns])
        return ws

    # ---------- 00_README ----------
    readme = []
    def rm(item, nilai, catatan=""):
        readme.append({"item": item, "nilai": nilai, "catatan": catatan})
    rm("toko", "Toko X", "nama disamarkan (BLOOMIE)")
    rm("rentang_tanggal", f"{min(r['date'] for r in rows).isoformat()} s/d {last_date.isoformat()}")
    rm("reopening", reopen.isoformat())
    for p in ("PRA", "PASCA", "PRA_BASELINE"):
        rm(f"hari_{p}", DAYS[p], "normalisasi per-30-hari = total/hari*30")
    for p in ("PRA", "PASCA"):
        sub = [r for r in rows if r["periode"] == p]
        rm(f"revenue_{p}", round(sum(r["rev"] for r in sub)))
        rm(f"nota_{p}", len({r["doc"] for r in sub}))
        rm(f"customer_unik_{p}", len({r["cust"] for r in sub if r["cust"]}))
    rm("target_brief1", "PRA 9.339.437.720/10.860 ; PASCA 6.622.181.617/9.767", "total harus SAMA (SKU induk tak ubah total)")
    # tanpa outlier
    sub_no = [r for r in rows if r["doc"] not in outlier_docs]
    rm("revenue_total_tanpa_outlier", round(sum(r["rev"] for r in sub_no)),
       f"outlier akun {OUTLIER_AKUN}, {len(outlier_docs)} nota dibuang untuk versi ini")
    rm("item_no_unik", n_item)
    rm("sku_induk_unik", n_sku, f"{n_item - n_sku} item_no digabung")
    rm("item_no_digabung", n_digabung)
    for pola in ("huruf_1-3", "angka_2digit", "huruf_nempel", "standar_4segmen"):
        rm(f"pola_{pola}", pola_counts.get(pola, 0))
    rm("item_perlu_review", n_review, "lihat sheet 01 perlu_review=TRUE")
    rm("baris_dibuang_retur_revneg", dropped_retur)
    rm("baris_dibuang_qty0", dropped_qty0)
    rm("baris_dibuang_tanpa_tanggal", dropped_nodate)
    rm("unit_kosong", unit_empty, "diisi '—'")
    rm("akun_walkin", ", ".join(sorted(WALKIN)), "is_walkin=TRUE; akun kasir & MEMBER (~72-79% nota)")
    rm("akun_outlier", OUTLIER_AKUN, "is_outlier=TRUE; nota 42.524 pcs 31 Mar 2026")
    rm("akun_generik_terdeteksi", len(generik))
    rm("customer_unik_total", len(cust_pra | cust_pas))
    rm("customer_overlap_PRA_PASCA", len(cust_pra & cust_pas))
    # ringkasan uji (diisi di bawah setelah hitung)
    README_SLOT = len(readme)
    sheet("00_README", ["item", "nilai", "catatan"], readme)

    # ---------- 01_SKU_Mapping ----------
    d01 = []
    for it in sorted(items_desc):
        m = mapping[it]
        d01.append({"item_no": it, "deskripsi": items_desc.get(it, ""), "sku_induk": m["sku_induk"],
                    "pola_suffix": m["pola"], "digabung": m["digabung"], "alasan": m["alasan"],
                    "perlu_review": m["perlu_review"]})
    sheet("01_SKU_Mapping", ["item_no","deskripsi","sku_induk","pola_suffix","digabung","alasan","perlu_review"], d01)

    # ---------- 02_Segmen_Nota ----------
    d02 = []
    for p in ("PRA", "PASCA", "PRA_BASELINE"):
        pred = (lambda n: n["baseline"]) if p == "PRA_BASELINE" else (lambda n, pp=p: n["periode"] == pp)
        for grp in ("AKUN", "WALKIN"):
            pn = [n for n in allnotas if pred(n) and kelompok_of(n) == grp]
            totn = len(pn); totr = sum(n["rev"] for n in pn) or 1
            for s in SEG_ORDER:
                sn = [n for n in pn if n["seg"] == s]
                cnt = len(sn); rev = sum(n["rev"] for n in sn); nilai = [n["rev"] for n in sn]
                d02.append({"periode": p, "kelompok": grp, "segmen": s, "label": SEG_LABEL[s],
                            "jumlah_nota": cnt, "nota_per30": round(per30(cnt, DAYS[p]), 2),
                            "pct_nota": round(cnt/totn, 4) if totn else "", "revenue": round(rev),
                            "revenue_per30": round(per30(rev, DAYS[p])), "pct_revenue": round(rev/totr, 4),
                            "aov": round(rev/cnt) if cnt else "", "median_nilai_nota": round(median(nilai)) if nilai else "",
                            "avg_jumlah_sku_induk": round(sum(n["n_sku"] for n in sn)/cnt, 2) if cnt else "",
                            "avg_qty_per_sku": round(sum(n["max_qty_sku"] for n in sn)/cnt, 2) if cnt else "",
                            "jumlah_akun_unik": len({n["cust"] for n in sn if n["cust"]})})
    sheet("02_Segmen_Nota", ["periode","kelompok","segmen","label","jumlah_nota","nota_per30","pct_nota","revenue",
                             "revenue_per30","pct_revenue","aov","median_nilai_nota","avg_jumlah_sku_induk",
                             "avg_qty_per_sku","jumlah_akun_unik"], d02)

    # ---------- 03_Lama_vs_Baru (tier lama x segmen baru) ----------
    d03 = []
    for p in ("PRA", "PASCA"):
        cell = defaultdict(lambda: {"n": 0, "rev": 0.0})
        for n in allnotas:
            if n["periode"] != p: continue
            e = cell[(n["tier_lama"], n["seg"])]; e["n"] += 1; e["rev"] += n["rev"]
        for t in TIER_ORDER:
            for s in SEG_ORDER:
                e = cell.get((t, s))
                if not e: continue
                d03.append({"periode": p, "tier_lama": t, "segmen_baru": s,
                            "jumlah_nota": e["n"], "revenue": round(e["rev"])})
    sheet("03_Lama_vs_Baru", ["periode","tier_lama","segmen_baru","jumlah_nota","revenue"], d03)

    # ---------- 04_Kategori_x_Segmen ----------
    # per (periode,kat,seg): revenue, dan metrik line (median qty/sku-line, % line kelipatan 6/12)
    kat_rev_p = defaultdict(float)                               # (periode,kat) -> rev total kat
    cell04 = defaultdict(lambda: {"rev": 0.0, "lineqty": []})    # (periode,kat,seg) -> ...
    for r in rows:
        n = notas[r["doc"]]
        kat_rev_p[(r["periode"], r["kat"])] += r["rev"]
    # line-level per (periode,kat,seg): kumpulkan qty per (doc,sku,unit)
    line_group = defaultdict(float)                             # (periode,kat,seg,doc,sku,unit) -> qty
    line_rev = defaultdict(float)
    for r in rows:
        n = notas[r["doc"]]
        k = (r["periode"], r["kat"], n["seg"], r["doc"], r["sku"], r["unit"])
        line_group[k] += r["qty"]; line_rev[k] += r["rev"]
    for (p, kat, seg, doc, sku, unit), q in line_group.items():
        e = cell04[(p, kat, seg)]; e["rev"] += line_rev[(p, kat, seg, doc, sku, unit)]; e["lineqty"].append(q)
    d04 = []
    for (p, kat, seg) in sorted(cell04, key=lambda x: (x[0], x[1], SEG_ORDER.index(x[2]))):
        e = cell04[(p, kat, seg)]; lq = e["lineqty"]; nl = len(lq) or 1
        d04.append({"periode": p, "kategori": kat, "segmen": seg, "revenue": round(e["rev"]),
                    "pct_revenue_kategori": round(e["rev"]/kat_rev_p[(p, kat)], 4) if kat_rev_p[(p, kat)] else "",
                    "median_qty_per_sku_line": round(median(lq), 2) if lq else "",
                    "pct_line_kelipatan_6": round(sum(1 for q in lq if q and q % 6 == 0)/nl, 4),
                    "pct_line_kelipatan_12": round(sum(1 for q in lq if q and q % 12 == 0)/nl, 4)})
    sheet("04_Kategori_x_Segmen", ["periode","kategori","segmen","revenue","pct_revenue_kategori",
                                   "median_qty_per_sku_line","pct_line_kelipatan_6","pct_line_kelipatan_12"], d04)

    # ---------- customer-period agregasi ----------
    cp = defaultdict(lambda: {"notas": 0, "grosir": 0, "rev": 0.0, "qty": 0.0,
                              "skus": set(), "lines": 0, "catrev": defaultdict(float), "maxs": []})
    for n in allnotas:
        if not n["cust"]: continue
        e = cp[(n["cust"], n["periode"])]
        e["notas"] += 1; e["rev"] += n["rev"]; e["qty"] += n["qty"]
        if n["grosir_sku"]: e["grosir"] += 1
        e["skus"] |= n["skus"]; e["lines"] += len(n["line"]); e["maxs"].append(n["max_qty_sku"])
        for k, v in n["catrev"].items(): e["catrev"][k] += v
    cust_first = {}; cust_last = {}; cust_name = {}; cust_months = defaultdict(set)
    for n in allnotas:
        c = n["cust"]
        if not c: continue
        cust_first[c] = n["date"] if c not in cust_first else min(cust_first[c], n["date"])
        cust_last[c] = n["date"] if c not in cust_last else max(cust_last[c], n["date"])
        cust_name.setdefault(c, n["name"]); cust_months[c].add(n["bulan"])

    def tipe_of(c, periode):
        if c in WALKIN: return "Walk-in"
        e = cp.get((c, periode))
        if not e or not e["notas"]: return ""
        pct = e["grosir"] / e["notas"]
        return "Grosir" if pct >= 0.8 else ("Campuran" if pct >= 0.2 else "Eceran")

    def topcat(e):
        return max(e["catrev"].items(), key=lambda x: x[1])[0] if e and e["catrev"] else ""

    # ---------- 05_Customer_Level ----------
    d05 = []
    for c in (cust_pra | cust_pas):
        P = cp.get((c, "PRA")); Q = cp.get((c, "PASCA"))
        status = "Kembali" if (c in cust_pra and c in cust_pas) else ("Belum kembali" if c in cust_pra else "Baru")
        rev_pra = P["rev"] if P else 0; rev_pas = Q["rev"] if Q else 0
        d05.append({
            "kode_customer": c, "nama_customer": cust_name.get(c, ""), "wilayah": wilayah_of(c),
            "kelompok": "WALKIN" if c in WALKIN else "AKUN",
            "is_walkin": c in WALKIN, "is_outlier": c == OUTLIER_AKUN, "is_generik": c in generik,
            "status": status,
            "kategori_utama_pra": topcat(P), "kategori_utama_pasca": topcat(Q),
            "nota_pra": P["notas"] if P else 0, "revenue_pra": round(rev_pra),
            "nota_pasca": Q["notas"] if Q else 0, "revenue_pasca": round(rev_pas),
            "revenue_per30_pra": round(per30(rev_pra, DAYS["PRA"])), "revenue_per30_pasca": round(per30(rev_pas, DAYS["PASCA"])),
            "jumlah_sku_induk_unik_pra": len(P["skus"]) if P else 0, "jumlah_sku_induk_unik_pasca": len(Q["skus"]) if Q else 0,
            "avg_qty_per_sku_pra": round(P["qty"]/P["lines"], 2) if P and P["lines"] else "",
            "avg_qty_per_sku_pasca": round(Q["qty"]/Q["lines"], 2) if Q and Q["lines"] else "",
            "pct_nota_grosir_sku_pra": round(P["grosir"]/P["notas"], 4) if P and P["notas"] else "",
            "pct_nota_grosir_sku_pasca": round(Q["grosir"]/Q["notas"], 4) if Q and Q["notas"] else "",
            "tipe_pembeli_pra": tipe_of(c, "PRA"), "tipe_pembeli_pasca": tipe_of(c, "PASCA"),
            "tgl_pertama_beli": cust_first[c].isoformat(), "tgl_terakhir_beli": cust_last[c].isoformat(),
            "bulan_pertama_beli": month_bucket(cust_first[c], reopen), "jumlah_bulan_aktif": len(cust_months[c]),
            "_sort": rev_pra + rev_pas,
        })
    d05.sort(key=lambda x: -x["_sort"])
    cols05 = ["kode_customer","nama_customer","wilayah","kelompok","is_walkin","is_outlier","is_generik","status",
              "kategori_utama_pra","kategori_utama_pasca","nota_pra","revenue_pra","nota_pasca","revenue_pasca",
              "revenue_per30_pra","revenue_per30_pasca","jumlah_sku_induk_unik_pra","jumlah_sku_induk_unik_pasca",
              "avg_qty_per_sku_pra","avg_qty_per_sku_pasca","pct_nota_grosir_sku_pra","pct_nota_grosir_sku_pasca",
              "tipe_pembeli_pra","tipe_pembeli_pasca","tgl_pertama_beli","tgl_terakhir_beli","bulan_pertama_beli",
              "jumlah_bulan_aktif"]
    sheet("05_Customer_Level", cols05, d05)

    # ---------- 06_Tipe_x_Status (akun non-walk-in) ----------
    d06map = defaultdict(lambda: {"akun": set(), "rp": 0.0, "rq": 0.0, "r30p": [], "r30q": []})
    for row in d05:
        c = row["kode_customer"]
        if c in WALKIN: continue
        # tipe pembeli: pakai pasca bila ada, else pra
        tipe = row["tipe_pembeli_pasca"] or row["tipe_pembeli_pra"] or "-"
        e = d06map[(tipe, row["status"])]
        e["akun"].add(c); e["rp"] += row["revenue_pra"]; e["rq"] += row["revenue_pasca"]
        e["r30p"].append(row["revenue_per30_pra"]); e["r30q"].append(row["revenue_per30_pasca"])
    d06 = []
    for (tipe, st), e in sorted(d06map.items()):
        d06.append({"tipe_pembeli": tipe, "status": st, "jumlah_akun": len(e["akun"]),
                    "revenue_pra": round(e["rp"]), "revenue_pasca": round(e["rq"]),
                    "rata2_rev_per30_pra": round(sum(e["r30p"])/len(e["r30p"])) if e["r30p"] else "",
                    "rata2_rev_per30_pasca": round(sum(e["r30q"])/len(e["r30q"])) if e["r30q"] else ""})
    sheet("06_Tipe_x_Status", ["tipe_pembeli","status","jumlah_akun","revenue_pra","revenue_pasca",
                               "rata2_rev_per30_pra","rata2_rev_per30_pasca"], d06)

    # ---------- 07_Migrasi_Tipe (akun Kembali) ----------
    d07map = defaultdict(lambda: {"n": 0, "rp": 0.0, "rq": 0.0})
    for row in d05:
        if row["status"] != "Kembali" or row["kode_customer"] in WALKIN: continue
        e = d07map[(row["tipe_pembeli_pra"] or "-", row["tipe_pembeli_pasca"] or "-")]
        e["n"] += 1; e["rp"] += row["revenue_pra"]; e["rq"] += row["revenue_pasca"]
    d07 = [{"tipe_pra": k[0], "tipe_pasca": k[1], "jumlah_akun": v["n"],
            "revenue_pra": round(v["rp"]), "revenue_pasca": round(v["rq"])} for k, v in sorted(d07map.items())]
    sheet("07_Migrasi_Tipe", ["tipe_pra","tipe_pasca","jumlah_akun","revenue_pra","revenue_pasca"], d07)

    # ---------- 08_Walkin_Dalam ----------
    def walkin_block(label, pn):
        totn = len(pn) or 1
        gn = [n for n in pn if n["grosir_sku"]]
        return {"scope": label, "jumlah_nota": len(pn), "pct_nota_grosir_sku": round(len(gn)/totn, 4),
                "revenue_nota_grosir_sku": round(sum(n["rev"] for n in gn)),
                "avg_jumlah_sku_induk": round(sum(n["n_sku"] for n in pn)/len(pn), 2) if pn else "",
                "avg_qty_per_sku": round(sum(n["max_qty_sku"] for n in pn)/len(pn), 2) if pn else ""}
    d08 = []
    for p in ("PRA", "PASCA"):
        d08.append(walkin_block(p, [n for n in allnotas if n["walkin"] and n["periode"] == p]))
    for c in sorted(WALKIN):
        for p in ("PRA", "PASCA"):
            d08.append(walkin_block(f"{c}·{p}", [n for n in allnotas if n["cust"] == c and n["periode"] == p]))
    # per minggu relatif reopening
    def monday(d): return d - timedelta(days=d.weekday())
    reopen_monday = monday(reopen)
    wk = defaultdict(list)
    for n in allnotas:
        if not n["walkin"]: continue
        rel = (monday(n["date"]) - reopen_monday).days // 7
        if rel < -12: continue
        wk[rel].append(n)
    for rel in sorted(wk):
        lab = f"minggu_{rel:+d} ({(reopen_monday+timedelta(days=rel*7)).isoformat()})"
        d08.append(walkin_block(lab, wk[rel]))
    sheet("08_Walkin_Dalam", ["scope","jumlah_nota","pct_nota_grosir_sku","revenue_nota_grosir_sku",
                              "avg_jumlah_sku_induk","avg_qty_per_sku"], d08)

    # ---------- 09_Retensi_per_Tipe ----------
    first_nota = {}
    for n in sorted(allnotas, key=lambda x: x["date"]):
        if n["cust"] and n["cust"] not in first_nota:
            first_nota[n["cust"]] = n
    cust_dates = defaultdict(list)
    for n in allnotas:
        if n["cust"]: cust_dates[n["cust"]].append(n["date"])

    def first_tipe(c):
        if c in WALKIN: return "Walk-in"
        return "Grosir" if first_nota[c]["grosir_sku"] else "Eceran"

    cust_month_first = {c: month_bucket(cust_first[c], reopen) for c in cust_first}
    months = sorted({r["bulan"] for r in rows}, key=mkey)
    d09map = defaultdict(lambda: {"members": [], "aktif": []})
    for c in cust_first:
        d09map[(cust_month_first[c], first_tipe(c))]["members"].append(c)
    d09 = []
    for (bln, tipe), e in sorted(d09map.items(), key=lambda x: (mkey(x[0][0]), x[0][1])):
        members = e["members"]; nmem = len(members) or 1
        row = {"bulan_pertama_beli": bln, "tipe_nota_pertama": tipe, "ukuran": len(members)}
        for w in (30, 60, 90):
            enough = (last_date - cust_first[members[0]]).days >= w
            if not enough:
                row[f"pct_{w}"] = "NA"; continue
            cnt = 0
            for c in members:
                f = first_nota[c]["date"]
                if any(f < d <= f + timedelta(days=w) for d in cust_dates[c]): cnt += 1
            row[f"pct_{w}"] = round(cnt/len(members), 4)
        row["avg_bulan_aktif"] = round(sum(len(cust_months[c]) for c in members)/nmem, 2)
        d09.append(row)
    sheet("09_Retensi_per_Tipe", ["bulan_pertama_beli","tipe_nota_pertama","ukuran","pct_30","pct_60","pct_90","avg_bulan_aktif"], d09)

    # ---------- 10_SKU_Level ----------
    sku_p = defaultdict(lambda: {"rev": 0.0, "qty": 0.0, "notas": set(), "qty_gros": 0.0, "custs": set()})
    sku_first = {}
    for r in rows:
        n = notas[r["doc"]]
        e = sku_p[(r["sku"], r["periode"])]
        e["rev"] += r["rev"]; e["qty"] += r["qty"]; e["notas"].add(r["doc"])
        if n["grosir_sku"]: e["qty_gros"] += r["qty"]
        if r["cust"]: e["custs"].add(r["cust"])
        sku_first[r["sku"]] = r["date"] if r["sku"] not in sku_first else min(sku_first[r["sku"]], r["date"])
    d10 = []
    for (sku, p), e in sorted(sku_p.items(), key=lambda x: -x[1]["rev"]):
        d10.append({"sku_induk": sku, "periode": p, "revenue": round(e["rev"]), "qty": round(e["qty"]),
                    "jumlah_nota": len(e["notas"]),
                    "pct_qty_dari_nota_grosir": round(e["qty_gros"]/e["qty"], 4) if e["qty"] else "",
                    "jumlah_akun_unik": len(e["custs"]), "pertama_kali_terjual": sku_first[sku].isoformat()})
    sheet("10_SKU_Level", ["sku_induk","periode","revenue","qty","jumlah_nota","pct_qty_dari_nota_grosir",
                           "jumlah_akun_unik","pertama_kali_terjual"], d10)

    # ---------- 11_SKU_Ringkas ----------
    sku_all = defaultdict(lambda: {"rev": 0.0, "qty": 0.0, "qty_gros": 0.0})
    for r in rows:
        n = notas[r["doc"]]
        e = sku_all[r["sku"]]; e["rev"] += r["rev"]; e["qty"] += r["qty"]
        if n["grosir_sku"]: e["qty_gros"] += r["qty"]
    sku_pra = {sku for (sku, p) in sku_p if p == "PRA"}
    sku_pas = {sku for (sku, p) in sku_p if p == "PASCA"}
    sku_baru = {sku for sku, d in sku_first.items() if d >= reopen}
    rev_total = sum(e["rev"] for e in sku_all.values()) or 1
    rev_baru = sum(sku_all[s]["rev"] for s in sku_baru)
    top20 = sorted(sku_all.values(), key=lambda e: -e["rev"])[:20]
    sku_gros = sum(1 for e in sku_all.values() if e["qty"] and e["qty_gros"]/e["qty"] >= 0.7)
    sku_ecer = sum(1 for e in sku_all.values() if e["qty"] and e["qty_gros"]/e["qty"] < 0.3)
    d11 = [
        {"metrik": "sku_induk_aktif_PRA", "nilai": len(sku_pra)},
        {"metrik": "sku_induk_aktif_PASCA", "nilai": len(sku_pas)},
        {"metrik": "sku_induk_baru_pasca_reopening", "nilai": len(sku_baru)},
        {"metrik": "share_revenue_sku_baru", "nilai": round(rev_baru/rev_total, 4)},
        {"metrik": "share_revenue_top20_sku", "nilai": round(sum(e["rev"] for e in top20)/rev_total, 4)},
        {"metrik": "sku_grosir_>=70pct_qty_dari_nota_grosir", "nilai": sku_gros},
        {"metrik": "sku_eceran_<30pct_qty_dari_nota_grosir", "nilai": sku_ecer},
        {"metrik": "sku_induk_unik_total", "nilai": len(sku_all)},
    ]
    sheet("11_SKU_Ringkas", ["metrik", "nilai"], d11)

    # ---------- 12_Sensitivitas ----------
    d12 = []
    for grp in ("AKUN", "WALKIN", "SEMUA"):
        for p in ("PRA", "PASCA"):
            pn = [n for n in allnotas if n["periode"] == p and (grp == "SEMUA" or kelompok_of(n) == grp)]
            totn = len(pn) or 1; totr = sum(n["rev"] for n in pn) or 1
            for thr in (6, 12, 24):
                gn = [n for n in pn if n["bulk"] or n["max_qty_sku"] >= thr]
                d12.append({"kelompok": grp, "periode": p, "ambang_qty_per_sku": thr,
                            "pct_nota_grosir": round(len(gn)/totn, 4),
                            "pct_revenue_grosir": round(sum(n["rev"] for n in gn)/totr, 4)})
    sheet("12_Sensitivitas", ["kelompok","periode","ambang_qty_per_sku","pct_nota_grosir","pct_revenue_grosir"], d12)

    # ---------- isi ringkasan uji ke README ----------
    def pct_seg_change():
        diff = sum(1 for n in allnotas if n["grosir_lama"] != n["grosir_sku"])
        return diff / len(allnotas) if allnotas else 0
    def pct_walkin_grosir():
        wn = [n for n in allnotas if n["walkin"]]
        return sum(1 for n in wn if n["grosir_sku"]) / len(wn) if wn else 0
    def pct_rev_grosir(p):
        pn = [n for n in allnotas if n["periode"] == p]; tr = sum(n["rev"] for n in pn) or 1
        return sum(n["rev"] for n in pn if n["grosir_sku"]) / tr
    uji = [
        ("uji_1_item_no_ke_sku", f"{n_item} item_no -> {n_sku} sku induk"),
        ("uji_2_pct_nota_pindah_klasifikasi", round(pct_seg_change(), 4)),
        ("uji_3_pct_nota_walkin_grosir_sku", round(pct_walkin_grosir(), 4)),
        ("uji_4_pct_revenue_grosir_sku_PRA", round(pct_rev_grosir("PRA"), 4)),
        ("uji_4_pct_revenue_grosir_sku_PASCA", round(pct_rev_grosir("PASCA"), 4)),
        ("uji_5_item_perlu_review", n_review),
    ]
    ws_readme = wb["00_README"]
    for item, nilai in uji:
        ws_readme.append([item, nilai, "ringkasan uji wajib"])

    # ================= SIMPAN =================
    B.autosize(wb, cap=34)
    os.makedirs(args.out, exist_ok=True)
    out = os.path.join(args.out, "TokoX_Analisis_SKU.xlsx")
    wb.save(out)

    # ---- validasi + ringkasan ke console ----
    print("\n=== VALIDASI ===")
    for p in ("PRA", "PASCA"):
        sub = [r for r in rows if r["periode"] == p]
        print(f"  {p}: revenue Rp{round(sum(r['rev'] for r in sub)):,} | nota {len({r['doc'] for r in sub})} | customer {len({r['cust'] for r in sub if r['cust']})}")
    print("  Target brief 1: PRA 9.339.437.720/10.860 | PASCA 6.622.181.617/9.767")
    print(f"  baris 01_SKU_Mapping = {len(d01)} (harus = item_no unik = {n_item})")
    seg_sum = Counter(n["seg"] for n in allnotas)
    print(f"  cek tiap nota tepat 1 segmen: sum={sum(seg_sum.values())} vs nota={len(allnotas)}")

    print("\n=== RINGKASAN UJI (tempel ke chat) ===")
    for item, nilai in uji:
        print(f"  {item}: {nilai}")
    print(f"\n  item_no unik -> sku induk: {n_item} -> {n_sku} ({n_digabung} digabung)")
    print(f"  pola: {dict(pola_counts)}")
    print(f"  % nota pindah klasifikasi (grosir lama vs sku): {pct_seg_change():.1%}")
    print(f"  % nota walk-in grosir-sku: {pct_walkin_grosir():.1%}")
    print(f"  % revenue grosir-sku: PRA {pct_rev_grosir('PRA'):.1%} -> PASCA {pct_rev_grosir('PASCA'):.1%}")
    # top review by revenue (maks 30)
    rev_item = defaultdict(float)
    for r in rows: rev_item[r["item"]] += r["rev"]
    review_items = sorted((it for it in items_desc if mapping[it]["perlu_review"]), key=lambda it: -rev_item[it])[:30]
    print(f"\n  TOP perlu_review (maks 30, by revenue): {len(review_items)} item")
    for it in review_items[:30]:
        print(f"    {it}  Rp{round(rev_item[it]):,}  | {items_desc.get(it,'')[:50]}")
    print(f"\n✓ Excel: {out}")


if __name__ == "__main__":
    main()
