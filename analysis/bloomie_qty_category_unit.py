"""
BLOOMIE — Tren jumlah pembelian (Quantity) per Kategori × Unit, per bulan (Juli pecah pra/pasca).
Mematuhi aturan: Quantity hanya dijumlah dalam SATUAN yang sama -> baris = (Kategori · Unit).
Kolom = periode (bulan; Juli dibelah pra/pasca) + Total. Header kolom Jul-pasca ditandai (reopening).

Output Excel ke folder di luar repo (default ~/Bloomie_Analysis/).
  python analysis/bloomie_qty_category_unit.py [--reopen 2026-07-19] [--out DIR]
"""
import os, argparse
from collections import defaultdict
import openpyxl
import _bloomie_common as B


def _qty(v):
    try:
        f = float(v)
        return int(f) if f == int(f) else f
    except Exception:
        return 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--reopen", default="2026-07-19")
    ap.add_argument("--out", default=os.path.join(os.path.expanduser("~"), "Bloomie_Analysis"))
    args = ap.parse_args()
    from datetime import datetime
    reopen = datetime.fromisoformat(args.reopen).date()

    rows = B.fetch_bloomie()
    print(f"{len(rows)} baris BLOOMIE.")
    if not rows:
        print("Tidak ada data BLOOMIE. Jalankan local_ingest.py dulu."); return

    qty = defaultdict(lambda: defaultdict(float))   # (kat,unit) -> periodkey -> qty
    docs = defaultdict(lambda: defaultdict(set))    # (kat,unit) -> periodkey -> set(nota)
    for r in rows:
        d = B.d_(r.get("posting_date"))
        if not d:
            continue
        kat = r.get("kategori") or "—"; unit = (r.get("unit") or "—").strip().upper()
        pk = B.period_key(d, reopen)
        qty[(kat, unit)][pk] += _qty(r.get("quantity"))
        if r.get("document_number") is not None:
            docs[(kat, unit)][pk].add(r.get("document_number"))

    periods = sorted({pk for m in qty.values() for pk in m}, key=B.period_order)
    combos = sorted(qty.keys(), key=lambda ku: -sum(qty[ku].values()))

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Qty per Kategori-Unit"
    ws.append(["BLOOMIE — Jumlah pembelian (Quantity) per Kategori × Unit, per bulan (reopening " + reopen.isoformat() + ")"])
    ws["A1"].font = openpyxl.styles.Font(bold=True, size=12)
    ws.append(["Quantity dijumlah hanya dalam satuan yang sama. Kolom pasca-reopening (Jul 19-akhir) ditandai kuning."])
    ws.append([])

    head = ["Kategori", "Unit"] + [B.period_label(k, reopen) for k in periods] + ["TOTAL"]
    ws.append(head)
    hrow = ws[ws.max_row]
    for c in hrow:
        c.font = B.HDR; c.fill = B.FILL
    # tandai kolom Jul-pasca (reopening) di header
    for idx, k in enumerate(periods):
        if k[1] == 1:
            hrow[2 + idx].fill = B.REOP; hrow[2 + idx].font = B.BOLD

    for (kat, unit) in combos:
        vals = [qty[(kat, unit)].get(k, 0) for k in periods]
        ws.append([kat, unit] + [round(v) for v in vals] + [round(sum(vals))])
        rr = ws[ws.max_row]
        for i in range(2, len(head)):
            rr[i].number_format = B.NUM
        # beri warna sel kolom pasca
        for idx, k in enumerate(periods):
            if k[1] == 1:
                rr[2 + idx].fill = B.REOP

    ws.append([])
    ws.append(["Catatan: TIDAK ada baris/kolom 'grand total' lintas satuan — menjumlah PCS + PAK + LSN tidak sah."])

    # ---------- Sheet 2: rata-rata quantity per order (nota) ----------
    ws2 = wb.create_sheet("Avg Qty per Order")
    ws2.append(["BLOOMIE — Rata-rata Quantity per order (nota) per Kategori × Unit"])
    ws2["A1"].font = openpyxl.styles.Font(bold=True, size=12)
    ws2.append(["Avg = total quantity ÷ jumlah nota yang memuat kategori-unit itu. 'tiap kali beli, rata-rata ambil berapa'."])
    ws2.append([])
    head2 = ["Kategori", "Unit"] + [B.period_label(k, reopen) for k in periods] + ["RATA-RATA KESELURUHAN"]
    ws2.append(head2)
    h2 = ws2[ws2.max_row]
    for c in h2:
        c.font = B.HDR; c.fill = B.FILL
    for idx, k in enumerate(periods):
        if k[1] == 1:
            h2[2 + idx].fill = B.REOP; h2[2 + idx].font = B.BOLD

    def avg(ku, k):
        n = len(docs[ku].get(k, set()))
        return round(qty[ku].get(k, 0) / n, 1) if n else None

    for (kat, unit) in combos:
        vals = [avg((kat, unit), k) for k in periods]
        tot_q = sum(qty[(kat, unit)].values())
        tot_d = len(set().union(*[docs[(kat, unit)].get(k, set()) for k in periods])) if periods else 0
        overall = round(tot_q / tot_d, 1) if tot_d else None
        ws2.append([kat, unit] + [("" if v is None else v) for v in vals] + [("" if overall is None else overall)])
        rr = ws2[ws2.max_row]
        for i in range(2, len(head2)):
            rr[i].number_format = '#,##0.0'
        for idx, k in enumerate(periods):
            if k[1] == 1:
                rr[2 + idx].fill = B.REOP
    ws2.append([])
    ws2.append(["Sel kosong = tak ada nota kategori-unit itu di periode tsb."])

    B.autosize(wb)
    os.makedirs(args.out, exist_ok=True)
    out = os.path.join(args.out, "bloomie_qty_kategori_unit.xlsx")
    wb.save(out)
    print("Periode:", ", ".join(B.period_label(k, reopen) for k in periods))
    print("✓ Excel:", out)


if __name__ == "__main__":
    main()
