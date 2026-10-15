"""
BLOOMIE — Tren komposisi customer: BARU vs RETAIN, per bulan (Juli pecah pra/pasca).
  Baru   = pembelian pertama customer (sepanjang data BLOOMIE) jatuh di periode ini.
  Retain = customer aktif di periode ini TAPI sudah pernah beli di periode sebelumnya.
Identitas customer = Customer/Vendor_Code (sesuai dashboard).

Output Excel ke folder di luar repo (default ~/Bloomie_Analysis/).
  python analysis/bloomie_customer_composition.py [--reopen 2026-07-19] [--out DIR]
"""
import os, argparse
from collections import defaultdict
import openpyxl
import _bloomie_common as B


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--reopen", default="2026-07-19")
    ap.add_argument("--out", default=os.path.join(os.path.expanduser("~"), "Bloomie_Analysis"))
    args = ap.parse_args()
    from datetime import datetime
    reopen = datetime.fromisoformat(args.reopen).date()

    rows = B.fetch_bloomie()
    print(f"{len(rows)} baris BLOOMIE.")
    if not rows:
        print("Tidak ada data BLOOMIE. Jalankan local_ingest.py dulu."); return

    # pembelian pertama tiap customer
    first = {}
    for r in rows:
        c = r.get("customer_code"); d = B.d_(r.get("posting_date"))
        if not c or not d:
            continue
        if c not in first or d < first[c]:
            first[c] = d
    first_pk = {c: B.period_key(d, reopen) for c, d in first.items()}

    # per periode: himpunan customer aktif + revenue per kelas
    per = defaultdict(lambda: {"baru": set(), "retain": set(), "rev_baru": 0.0, "rev_retain": 0.0})
    for r in rows:
        c = r.get("customer_code"); d = B.d_(r.get("posting_date"))
        if not c or not d:
            continue
        pk = B.period_key(d, reopen); rev = float(r.get(B.REV) or 0)
        kelas = "baru" if first_pk.get(c) == pk else "retain"
        per[pk][kelas].add(c)
        per[pk]["rev_" + kelas] += rev

    periods = sorted(per.keys(), key=B.period_order)

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Komposisi Customer"
    ws.append(["BLOOMIE — Komposisi customer Baru vs Retain per bulan (reopening " + reopen.isoformat() + ")"])
    ws["A1"].font = openpyxl.styles.Font(bold=True, size=13)
    ws.append([])
    head = ["Periode", "Customer Baru", "Customer Retain", "Total Customer",
            "% Baru", "% Retain", "Revenue Baru", "Revenue Retain"]
    ws.append(head)
    for c in ws[ws.max_row]:
        c.font = B.HDR; c.fill = B.FILL

    def add(label, nb, nr, rb, rr):
        tot = nb + nr
        ws.append([label, nb, nr, tot,
                   f"{nb/tot*100:.1f}%" if tot else "—", f"{nr/tot*100:.1f}%" if tot else "—",
                   round(rb), round(rr)])
        row = ws[ws.max_row]
        for i in (1, 2, 3): row[i].number_format = B.NUM
        for i in (6, 7): row[i].number_format = B.NUM

    prev_half = None
    tot_nb = tot_nr = 0; tot_rb = tot_rr = 0.0
    for k in periods:
        if k[1] == 1 and prev_half == 0:
            ws.append([f"———  REOPENING {reopen.isoformat()}  ———"])
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=len(head))
            cc = ws.cell(ws.max_row, 1); cc.fill = B.REOP; cc.font = B.BOLD; cc.alignment = B.CENTER
        p = per[k]
        add(B.period_label(k, reopen), len(p["baru"]), len(p["retain"]), p["rev_baru"], p["rev_retain"])
        prev_half = k[1]

    # catatan: total customer unik lintas periode bukan penjumlahan kolom (seorang bisa aktif di banyak bulan)
    ws.append([])
    uniq = len({c for c in first})
    ws.append([f"Catatan: total customer unik BLOOMIE = {uniq}. Angka per baris adalah customer AKTIF di bulan itu; "
               "satu customer bisa muncul di beberapa bulan, jadi kolom tidak untuk dijumlahkan."])

    B.autosize(wb)
    os.makedirs(args.out, exist_ok=True)
    out = os.path.join(args.out, "bloomie_komposisi_customer.xlsx")
    wb.save(out)
    print("Periode:", ", ".join(B.period_label(k, reopen) for k in periods))
    print("✓ Excel:", out)


if __name__ == "__main__":
    main()
