"""
Toko X (BLOOMIE) — Analisis SKU dari 119 akun WIN-BACK (C4_Winback DEFAULT + Perbaikan 1).
Memakai mapping sku_induk DEFAULT (Brief 2) & aturan Perbaikan 1 (Brief 3).

Hitung:
  1. sku_induk UNIK yang dibeli 119 akun ini di PRA (tanpa double count) + revenue + qty.
  2. Dari SKU poin 1: berapa masih terjual di PASCA (oleh siapa pun) vs tidak terjual sama sekali.
  3. Dari SKU poin 1: berapa yang HANYA dibeli akun-akun ini di PRA (tak ada pembeli PRA lain)
     DAN tidak dibeli siapa pun di PASCA (SKU "yatim" yang tergantung pada akun hilang ini).
  4. Pecah poin 1-3 per tipe akun (Grosir/Campuran) & per kategori.
  5. Daftar SKU: sku_induk, deskripsi, kategori, jumlah akun win-back, qty & revenue PRA, status PASCA.

Output: ~/Bloomie_Analysis/TokoX_Winback_SKU.xlsx
  python analysis/tokox_winback_sku.py [--reopen 2026-07-19] [--out DIR]
"""
import os, argparse
from collections import defaultdict, Counter
from datetime import datetime, date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

import _bloomie_common as B
from tokox_analisis_sku import build_sku_mapping, CONV_PCS, seg_v2, WALKIN

REV = B.REV
BULK_KEMASAN = {"BOX", "PAK", "PACK", "DUS", "CTN", "KARTON", "BALL", "RIM", "KG"}


def qf(v):
    try: return float(v)
    except Exception: return 0.0


def eff(unit, qty):
    return ("PCS", qty * CONV_PCS[unit]) if unit in CONV_PCS else (unit, qty)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--reopen", default="2026-07-19")
    ap.add_argument("--out", default=os.path.join(os.path.expanduser("~"), "Bloomie_Analysis"))
    args = ap.parse_args()
    reopen = datetime.fromisoformat(args.reopen).date()
    jun1 = date(reopen.year, 6, 1)

    raw = B.fetch_bloomie()
    print(f"{len(raw)} baris mentah Toko X.")
    if not raw:
        print("Tidak ada data. Jalankan local_ingest.py dulu."); return

    # ---- cleaning (sama Brief 2/3) ----
    rows = []; desc_count = defaultdict(Counter)
    for r in raw:
        d = B.d_(r.get("posting_date"))
        if not d: continue
        q = qf(r.get("quantity")); rev = qf(r.get(REV))
        if rev < 0 or q <= 0: continue
        u = (r.get("unit") or "").strip().upper()
        if u == "PIECES": u = "PCS"
        if u == "": u = "—"
        item = (r.get("item_no") or "").strip()
        dsc = (r.get("item_description") or "").strip()
        if dsc: desc_count[item][dsc] += 1
        rows.append({"doc": r.get("document_number"), "date": d,
                     "cust": (r.get("customer_code") or "").strip(), "item": item,
                     "kat": (r.get("kategori") or "Lainnya"), "qty": q, "unit": u, "rev": rev})

    items_desc = {it: c.most_common(1)[0][0] if c else "" for it, c in desc_count.items()}
    for r in rows: items_desc.setdefault(r["item"], "")
    mapping = build_sku_mapping(items_desc)
    for r in rows:
        r["sku"] = mapping[r["item"]]["sku_induk"]
        r["periode"] = "PRA" if r["date"] < reopen else "PASCA"

    # ---- nota-level (Perbaikan 1: lusinan->pcs, kemasan BUKAN auto-grosir) ----
    notas = {}
    for r in rows:
        n = notas.get(r["doc"])
        if n is None:
            n = notas[r["doc"]] = {"date": r["date"], "cust": r["cust"], "line": defaultdict(float)}
        elif r["date"] < n["date"]:
            n["date"] = r["date"]
        u, q = eff(r["unit"], r["qty"])
        n["line"][(r["sku"], u)] += q
    for n in notas.values():
        n["periode"] = "PRA" if n["date"] < reopen else "PASCA"
        mx = max(n["line"].values()) if n["line"] else 0
        n["grosir"] = seg_v2(mx) in ("S4", "S5")

    # ---- tipe akun PRA + winback ----
    cp = defaultdict(lambda: {"n": 0, "g": 0})
    cust_last = {}; cust_per = defaultdict(set)
    for n in notas.values():
        c = n["cust"]
        if not c: continue
        cust_per[c].add(n["periode"])
        cust_last[c] = n["date"] if c not in cust_last else max(cust_last[c], n["date"])
        if n["periode"] == "PRA":
            e = cp[c]; e["n"] += 1
            if n["grosir"]: e["g"] += 1

    def tipe_pra(c):
        e = cp.get(c)
        if not e or not e["n"]: return ""
        r = e["g"] / e["n"]
        return "Grosir" if r >= 0.8 else ("Campuran" if r >= 0.2 else "Eceran")

    winback = {}   # cust -> tipe
    for c in cp:
        if c in WALKIN: continue
        if "PASCA" in cust_per[c]: continue              # bukan 'belum kembali'
        if cust_last.get(c, date.min) < jun1: continue
        tp = tipe_pra(c)
        if tp in ("Grosir", "Campuran"):
            winback[c] = tp
    print(f"Win-back terdeteksi: {len(winback)} akun (target 119).")

    # ---- maps SKU ----
    pra_buyers = defaultdict(set)      # sku -> semua cust PRA
    pasca_buyers = defaultdict(set)    # sku -> semua cust PASCA (incl walk-in)
    wb_sku = defaultdict(lambda: {"rev": 0.0, "qty": 0.0, "accts": set(), "kat": Counter()})
    wb_tipe = defaultdict(lambda: {"rev": 0.0, "qty": 0.0, "skus": set(), "accts": set()})
    for r in rows:
        sku = r["sku"]; c = r["cust"]
        if r["periode"] == "PRA":
            pra_buyers[sku].add(c)
            if c in winback:
                _, qe = eff(r["unit"], r["qty"])
                e = wb_sku[sku]; e["rev"] += r["rev"]; e["qty"] += qe
                e["accts"].add(c); e["kat"][r["kat"]] += r["rev"]
                t = wb_tipe[winback[c]]; t["rev"] += r["rev"]; t["qty"] += qe
                t["skus"].add(sku); t["accts"].add(c)
        else:
            pasca_buyers[sku].add(c)

    sku_set = set(wb_sku)                                   # poin 1
    def sold_pasca(sku): return len(pasca_buyers.get(sku, set())) > 0
    def orphan(sku):                                        # poin 3
        return pra_buyers.get(sku, set()) <= set(winback) and not pasca_buyers.get(sku, set())
    sku_kat = {s: (wb_sku[s]["kat"].most_common(1)[0][0] if wb_sku[s]["kat"] else "Lainnya") for s in sku_set}

    tot_rev = sum(e["rev"] for e in wb_sku.values())
    tot_qty = sum(e["qty"] for e in wb_sku.values())
    n_sold = sum(1 for s in sku_set if sold_pasca(s))
    n_notsold = len(sku_set) - n_sold
    n_orphan = sum(1 for s in sku_set if orphan(s))

    # ================= WORKBOOK =================
    HDR = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="C00000"); BOLD = Font(bold=True)
    wb = openpyxl.Workbook(); wb.remove(wb.active)

    def sheet(name, cols, data):
        ws = wb.create_sheet(name[:31]); ws.append(cols)
        for c in ws[1]: c.font = HDR; c.fill = FILL
        for row in data: ws.append([row.get(c, "") for c in cols])
        for i, c in enumerate(cols):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width = max(12, min(40, len(c)+4))
        return ws

    NUM = '#,##0'
    # 00_README / ringkasan
    d0 = [
        {"item": "toko", "nilai": "Toko X", "catatan": "BLOOMIE (disamarkan)"},
        {"item": "reopening", "nilai": reopen.isoformat(), "catatan": ""},
        {"item": "akun_winback", "nilai": len(winback), "catatan": "DEFAULT+Perbaikan1; target 119"},
        {"item": "P1_sku_unik_PRA", "nilai": len(sku_set), "catatan": "tanpa double count"},
        {"item": "P1_revenue_PRA", "nilai": round(tot_rev), "catatan": "revenue winback atas SKU ini di PRA"},
        {"item": "P1_qty_PRA_pcs_equiv", "nilai": round(tot_qty), "catatan": "lusinan dikonversi ke pcs"},
        {"item": "P2_sku_masih_terjual_PASCA", "nilai": n_sold, "catatan": "dibeli siapa pun di PASCA"},
        {"item": "P2_sku_TIDAK_terjual_PASCA", "nilai": n_notsold, "catatan": ""},
        {"item": "P3_sku_yatim", "nilai": n_orphan,
         "catatan": "hanya dibeli akun winback di PRA & tak ada pembeli di PASCA"},
        {"item": "P3_revenue_yatim_PRA", "nilai": round(sum(wb_sku[s]["rev"] for s in sku_set if orphan(s))), "catatan": ""},
    ]
    sheet("00_Ringkasan", ["item", "nilai", "catatan"], d0)

    # 02_Per_Tipe (SKU bisa overlap antar tipe bila dibeli akun kedua tipe)
    d2 = []
    for tipe in ("Grosir", "Campuran"):
        t = wb_tipe[tipe]; skus = t["skus"]
        d2.append({"tipe_akun": tipe, "jumlah_akun": len(t["accts"]), "jumlah_sku_unik": len(skus),
                   "revenue_pra": round(t["rev"]), "qty_pra_pcs": round(t["qty"]),
                   "sku_masih_terjual_pasca": sum(1 for s in skus if sold_pasca(s)),
                   "sku_tidak_terjual_pasca": sum(1 for s in skus if not sold_pasca(s)),
                   "sku_yatim": sum(1 for s in skus if orphan(s))})
    sheet("02_Per_Tipe", ["tipe_akun", "jumlah_akun", "jumlah_sku_unik", "revenue_pra", "qty_pra_pcs",
                          "sku_masih_terjual_pasca", "sku_tidak_terjual_pasca", "sku_yatim"], d2)

    # 03_Per_Kategori
    kat_sku = defaultdict(set)
    for s in sku_set: kat_sku[sku_kat[s]].add(s)
    d3 = []
    for kat, skus in sorted(kat_sku.items(), key=lambda kv: -sum(wb_sku[s]["rev"] for s in kv[1])):
        d3.append({"kategori": kat, "jumlah_sku_unik": len(skus),
                   "revenue_pra": round(sum(wb_sku[s]["rev"] for s in skus)),
                   "qty_pra_pcs": round(sum(wb_sku[s]["qty"] for s in skus)),
                   "sku_masih_terjual_pasca": sum(1 for s in skus if sold_pasca(s)),
                   "sku_tidak_terjual_pasca": sum(1 for s in skus if not sold_pasca(s)),
                   "sku_yatim": sum(1 for s in skus if orphan(s))})
    sheet("03_Per_Kategori", ["kategori", "jumlah_sku_unik", "revenue_pra", "qty_pra_pcs",
                              "sku_masih_terjual_pasca", "sku_tidak_terjual_pasca", "sku_yatim"], d3)

    # 04_Daftar_SKU
    d4 = []
    for s in sorted(sku_set, key=lambda x: -wb_sku[x]["rev"]):
        e = wb_sku[s]
        d4.append({"sku_induk": s, "deskripsi": items_desc.get(s, "") or "",
                   "kategori": sku_kat[s], "jumlah_akun_winback": len(e["accts"]),
                   "qty_pra_pcs": round(e["qty"]), "revenue_pra": round(e["rev"]),
                   "status_pasca": "Masih terjual" if sold_pasca(s) else "Tidak terjual",
                   "sku_yatim": orphan(s)})
    ws4 = sheet("04_Daftar_SKU", ["sku_induk", "deskripsi", "kategori", "jumlah_akun_winback",
                                  "qty_pra_pcs", "revenue_pra", "status_pasca", "sku_yatim"], d4)
    ws4.auto_filter.ref = f"A1:H{ws4.max_row}"
    ws4.freeze_panes = "A2"

    os.makedirs(args.out, exist_ok=True)
    out = os.path.join(args.out, "TokoX_Winback_SKU.xlsx")
    wb.save(out)

    print("\n=== RINGKASAN ===")
    print(f"  Win-back: {len(winback)} akun | SKU unik PRA: {len(sku_set)} | rev Rp{round(tot_rev):,} | qty {round(tot_qty):,} pcs")
    print(f"  PASCA: masih terjual {n_sold} | tidak terjual {n_notsold} | yatim {n_orphan}")
    print(f"\n✓ Excel: {out}")


if __name__ == "__main__":
    main()
