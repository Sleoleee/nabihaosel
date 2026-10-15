"""Helper bersama untuk report BLOOMIE (baca SQLite lokal, pembagian periode pra/pasca)."""
import os, sys
from datetime import datetime

_HERE = os.path.dirname(os.path.abspath(__file__))
_BACKEND = os.path.join(os.path.dirname(_HERE), "backend")
sys.path.insert(0, _BACKEND)
try:
    from dotenv import load_dotenv
    load_dotenv(os.path.join(_BACKEND, ".env"))
except Exception:
    pass

import config
import local_store
from openpyxl.styles import Font, PatternFill, Alignment

MONTHS = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]
REV = config.REVENUE_COLUMN

# gaya
HDR = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="D31137")
REOP = PatternFill("solid", fgColor="FFF3B0"); BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center")
NUM = '#,##0'


def d_(v):
    try:
        return datetime.fromisoformat(str(v)[:10]).date()
    except Exception:
        return None


def fetch_bloomie():
    conn = local_store.connect(); local_store.ensure_schema(conn)
    cw = local_store.channel_where("BLOOMIE")
    rows = [r for r in local_store.iter_rows(conn, where=cw[0], params=cw[1])
            if config.branch_group(r.get("branch")) == "BLOOMIE"]
    conn.close()
    return rows


def period_key(d, reopen):
    m = d.month
    if m == reopen.month and d.year == reopen.year:
        return (m, 0) if d.day < reopen.day else (m, 1)
    return (m, -1)


def period_order(k):
    m, half = k
    return (m, 0 if half == -1 else (1 if half == 0 else 2))


def period_label(k, reopen):
    m, half = k
    if half == -1:
        return MONTHS[m - 1]
    if half == 0:
        return f"{MONTHS[m-1]} (1-{reopen.day-1}) · pra"
    return f"{MONTHS[m-1]} ({reopen.day}-akhir) · pasca"


def autosize(wb, cap=26):
    for w in wb.worksheets:
        for col in w.columns:
            try:
                width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                w.column_dimensions[col[0].column_letter].width = min(cap, max(11, width + 2))
            except Exception:
                pass
