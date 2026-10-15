"""
BLOOMIE — Behaviour Scorecard: PRA vs PASCA reopening + indikasi RETAIL vs GROSIR.
Patokan grosir = GABUNGAN: nota memakai satuan bulk ATAU total qty/nota >= ambang.
Kesimpulan 'condong' berdasarkan share REVENUE dari nota grosir tiap sisi.

  python analysis/bloomie_behaviour_scorecard.py [--reopen 2026-07-19]
      [--qty-threshold 12] [--bulk "LSN,KODI,GROS,PAK,BOX,DUS,BALL,RIM,KARTON,LUSIN"] [--out DIR]
Output Excel ke folder di luar repo (default ~/Bloomie_Analysis/).
"""
import os, argparse
from collections import defaultdict
from datetime import datetime
import openpyxl
from openpyxl.styles import Font
import _bloomie_common as B

DEFAULT_BULK = "LSN,LUSIN,KODI,GROS,GROSS,PAK,PACK,BOX,DUS,BALL,RIM,KARTON,CTN"


def qf(v):
    try:
        return float(v)
    except Exception:
        return 0.0


def analyse(rows, reopen, bulk_units, qth):
    """Kembalikan metrik untuk satu sisi (sekumpulan baris)."""
    # agregasi per nota
    nota = defaultdict(lambda: {"qty": 0.0, "rev": 0.0, "bulk": False, "cust": None})
    cust_rev = defaultdict(float); cust_nota = defaultdict(set)
    unit_qty = defaultdict(float)
    for r in rows:
        doc = r.get("document_number")
        if doc is None:
            continue
        q = qf(r.get("quantity")); rev = qf(r.get(B.REV))
        unit = (r.get("unit") or "—").strip().upper()
        n = nota[doc]
        n["qty"] += q; n["rev"] += rev; n["cust"] = r.get("customer_code")
        if unit in bulk_units:
            n["bulk"] = True
        unit_qty[unit] += q
        c = r.get("customer_code")
        if c:
            cust_rev[c] += rev; cust_nota[c].add(doc)

    n_nota = len(nota)
    rev_tot = sum(v["rev"] for v in nota.values())
    custs = set(cust_rev)
    n_cust = len(custs)
    # jumlah bulan kalender yang ada di sisi ini (untuk normalisasi per-bulan)
    month_buckets = defaultdict(lambda: {"rev": 0.0, "cust": set(), "nota": set()})
    for r in rows:
        d = B.d_(r.get("posting_date"))
        if not d:
            continue
        mb = month_buckets[(d.year, d.month)]
        mb["rev"] += qf(r.get(B.REV))
        if r.get("customer_code"): mb["cust"].add(r.get("customer_code"))
        if r.get("document_number") is not None: mb["nota"].add(r.get("document_number"))
    n_months = max(1, len(month_buckets))
    avg_cust_month = sum(len(mb["cust"]) for mb in month_buckets.values()) / n_months
    rev_month = rev_tot / n_months
    nota_month = n_nota / n_months
    # grosir?
    def is_gros(v):
        return v["bulk"] or v["qty"] >= qth
    gros_nota = [v for v in nota.values() if is_gros(v)]
    gros_rev = sum(v["rev"] for v in gros_nota)
    # one-time / repeat
    onetime = sum(1 for c in custs if len(cust_nota[c]) <= 1)
    repeat = n_cust - onetime
    # konsentrasi top-5 customer
    top5 = sorted(cust_rev.values(), reverse=True)[:5]
    conc5 = (sum(top5) / rev_tot * 100) if rev_tot else 0
    # distribusi ukuran order (qty per nota)
    buckets = {"1": 0, "2-5": 0, "6-20": 0, "21-50": 0, ">50": 0}
    buck_rev = {k: 0.0 for k in buckets}
    for v in nota.values():
        q = v["qty"]
        k = "1" if q <= 1 else "2-5" if q <= 5 else "6-20" if q <= 20 else "21-50" if q <= 50 else ">50"
        buckets[k] += 1; buck_rev[k] += v["rev"]
    bulk_qty = sum(q for u, q in unit_qty.items() if u in bulk_units)
    all_qty = sum(unit_qty.values())
    return {
        "revenue": rev_tot, "nota": n_nota, "cust": n_cust,
        "n_months": n_months, "rev_month": rev_month, "nota_month": nota_month,
        "cust_month": avg_cust_month,
        "spend_per_cust_month": rev_month / avg_cust_month if avg_cust_month else 0,
        "aov": rev_tot / n_nota if n_nota else 0,
        "qty_per_nota": sum(v["qty"] for v in nota.values()) / n_nota if n_nota else 0,
        "spend_per_cust": rev_tot / n_cust if n_cust else 0,
        "onetime_pct": onetime / n_cust * 100 if n_cust else 0,
        "repeat_pct": repeat / n_cust * 100 if n_cust else 0,
        "conc5": conc5,
        "gros_nota_pct": len(gros_nota) / n_nota * 100 if n_nota else 0,
        "gros_rev_pct": gros_rev / rev_tot * 100 if rev_tot else 0,
        "bulk_qty_pct": bulk_qty / all_qty * 100 if all_qty else 0,
        "buckets": buckets, "buck_rev": buck_rev, "unit_qty": unit_qty,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--reopen", default="2026-07-19")
    ap.add_argument("--qty-threshold", type=float, default=12)
    ap.add_argument("--bulk", default=DEFAULT_BULK)
    ap.add_argument("--out", default=os.path.join(os.path.expanduser("~"), "Bloomie_Analysis"))
    args = ap.parse_args()
    reopen = datetime.fromisoformat(args.reopen).date()
    bulk_units = {u.strip().upper() for u in args.bulk.split(",") if u.strip()}
    qth = args.qty_threshold

    rows = B.fetch_bloomie()
    print(f"{len(rows)} baris BLOOMIE.")
    if not rows:
        print("Tidak ada data. Jalankan local_ingest.py dulu."); return
    pra_rows = [r for r in rows if (B.d_(r.get("posting_date")) and B.d_(r.get("posting_date")) < reopen)]
    pas_rows = [r for r in rows if (B.d_(r.get("posting_date")) and B.d_(r.get("posting_date")) >= reopen)]
    P = analyse(pra_rows, reopen, bulk_units, qth)
    Q = analyse(pas_rows, reopen, bulk_units, qth)

    wb = openpyxl.Workbook()

    # ---- Sheet 1: Scorecard ----
    ws = wb.active; ws.title = "Scorecard"
    ws.append([f"BLOOMIE — Behaviour Scorecard PRA vs PASCA (reopening {reopen.isoformat()})"]); ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Grosir = nota pakai satuan bulk ({', '.join(sorted(bulk_units))}) ATAU qty/nota >= {qth:g}"])
    ws.append([f"PRA = {P['n_months']} bulan · PASCA = {Q['n_months']} bulan. Metrik VOLUME dinormalisasi PER BULAN agar adil; metrik RASIO memang tak terpengaruh durasi."])
    ws.append([])
    ws.append(["Metrik", "PRA", "PASCA", "Δ (abs)", "Δ %", "Sinyal"])
    for c in ws[ws.max_row]: c.font = B.HDR; c.fill = B.FILL

    def chg(a, b): return f"{(b-a)/a*100:+.1f}%" if a else "—"
    def row(label, a, b, fmt="num", signal=""):
        if fmt == "num":
            va, vb, dv = round(a), round(b), round(b - a)
        elif fmt == "pct":
            va, vb, dv = f"{a:.1f}%", f"{b:.1f}%", f"{b-a:+.1f} poin"
        else:
            va, vb, dv = round(a, 1), round(b, 1), round(b - a, 1)
        ws.append([label, va, vb, dv, chg(a, b), signal])
    ws.append(["— VOLUME (per bulan) —"]); ws[ws.max_row][0].font = B.BOLD
    row("Revenue / bulan", P["rev_month"], Q["rev_month"])
    row("Nota / bulan", P["nota_month"], Q["nota_month"])
    row("Customer unik / bulan (rata2)", P["cust_month"], Q["cust_month"], "dec")
    row("Spending / customer / bulan", P["spend_per_cust_month"], Q["spend_per_cust_month"], signal="tinggi→grosir")
    ws.append(["— RASIO / INTENSITAS (tak terpengaruh durasi) —"]); ws[ws.max_row][0].font = B.BOLD
    row("AOV (revenue/nota)", P["aov"], Q["aov"], signal="tinggi→grosir")
    row("Qty per nota (kasar)", P["qty_per_nota"], Q["qty_per_nota"], "dec", "tinggi→grosir")
    row("% one-time buyer *", P["onetime_pct"], Q["onetime_pct"], "pct", "tinggi→retail")
    row("% repeat buyer *", P["repeat_pct"], Q["repeat_pct"], "pct", "tinggi→grosir")
    row("Konsentrasi top-5 cust", P["conc5"], Q["conc5"], "pct", "tinggi→grosir")
    row("% satuan bulk (dari qty)", P["bulk_qty_pct"], Q["bulk_qty_pct"], "pct", "tinggi→grosir")
    row("% nota grosir", P["gros_nota_pct"], Q["gros_nota_pct"], "pct", "tinggi→grosir")
    row("% revenue dari nota grosir", P["gros_rev_pct"], Q["gros_rev_pct"], "pct", "tinggi→grosir")
    ws.append([])
    ws.append(["* one-time/repeat sedikit bias: jendela PASCA lebih pendek → peluang beli-ulang lebih kecil. Baca sebagai indikatif."])
    ws.append([f"Total mentah (konteks saja): Revenue PRA {round(P['revenue']):,} / PASCA {round(Q['revenue']):,}; Nota PRA {P['nota']} / PASCA {Q['nota']}."])

    # ---- Sheet 2: Distribusi ukuran order ----
    ws2 = wb.create_sheet("Distribusi Order")
    ws2.append(["Distribusi ukuran order (qty per nota). Bandingkan pakai kolom %Nota atau Nota/bulan (adil lintas durasi)."]); ws2["A1"].font = Font(bold=True, size=12)
    ws2.append([])
    ws2.append(["Qty per nota", "Nota/bln PRA", "Nota/bln PASCA", "%Nota PRA", "%Nota PASCA", "Nota PRA (mentah)", "Nota PASCA (mentah)"])
    for c in ws2[ws2.max_row]: c.font = B.HDR; c.fill = B.FILL
    order = ["1", "2-5", "6-20", "21-50", ">50"]
    tnp, tnq = sum(P["buckets"].values()), sum(Q["buckets"].values())
    for k in order:
        ws2.append([k,
                    round(P["buckets"][k] / P["n_months"], 1), round(Q["buckets"][k] / Q["n_months"], 1),
                    f"{P['buckets'][k]/tnp*100:.1f}%" if tnp else "—", f"{Q['buckets'][k]/tnq*100:.1f}%" if tnq else "—",
                    P["buckets"][k], Q["buckets"][k]])

    # ---- Sheet 3: Sinyal satuan ----
    ws3 = wb.create_sheet("Sinyal Satuan")
    ws3.append(["Quantity per satuan (Unit) — pra vs pasca. Kolom 'Bulk?' = masuk daftar grosir."]); ws3["A1"].font = Font(bold=True, size=12)
    ws3.append(["Jika ada satuan yang salah klasifikasi, jalankan ulang dengan --bulk \"DAFTAR,BARU\"."])
    ws3.append([])
    ws3.append(["Unit", "Bulk?", "Qty/bln PRA", "Qty/bln PASCA", "Qty PRA (mentah)", "Qty PASCA (mentah)"])
    for c in ws3[ws3.max_row]: c.font = B.HDR; c.fill = B.FILL
    units = sorted(set(list(P["unit_qty"].keys()) + list(Q["unit_qty"].keys())),
                   key=lambda u: -(P["unit_qty"].get(u, 0) + Q["unit_qty"].get(u, 0)))
    for u in units:
        ws3.append([u, "ya" if u in bulk_units else "—",
                    round(P["unit_qty"].get(u, 0) / P["n_months"], 1), round(Q["unit_qty"].get(u, 0) / Q["n_months"], 1),
                    round(P["unit_qty"].get(u, 0)), round(Q["unit_qty"].get(u, 0))])

    # ---- Sheet 4: Indikasi retail/grosir (lean) ----
    ws4 = wb.create_sheet("Indikasi Retail-Grosir")
    ws4.append(["Indikasi condong RETAIL vs GROSIR"]); ws4["A1"].font = Font(bold=True, size=13)
    ws4.append([])

    def lean(m):
        return "condong GROSIR" if m["gros_rev_pct"] >= 50 else "condong RETAIL"
    ws4.append(["Dasar utama: % revenue dari nota grosir (>=50% = grosir)."])
    ws4.append(["Sisi", "% revenue grosir", "% nota grosir", "% one-time", "Spending/cust/bln", "KESIMPULAN"])
    for c in ws4[ws4.max_row]: c.font = B.HDR; c.fill = B.FILL
    ws4.append(["PRA", f"{P['gros_rev_pct']:.1f}%", f"{P['gros_nota_pct']:.1f}%", f"{P['onetime_pct']:.1f}%", round(P['spend_per_cust_month']), lean(P)])
    ws4.append(["PASCA", f"{Q['gros_rev_pct']:.1f}%", f"{Q['gros_nota_pct']:.1f}%", f"{Q['onetime_pct']:.1f}%", round(Q['spend_per_cust_month']), lean(Q)])
    ws4.append([])
    shift = "TIDAK berubah" if lean(P) == lean(Q) else f"BERGESER dari {lean(P).split()[-1]} → {lean(Q).split()[-1]}"
    ws4.append([f"Pergeseran pra→pasca: {shift}."])
    ws4.append([f"Catatan: jendela pasca pendek ({reopen.isoformat()}–akhir data) & efek reopening → indikatif."])
    for cell in (ws4[ws4.max_row-1][0], ws4[ws4.max_row][0]):
        cell.font = Font(italic=True)

    B.autosize(wb)
    os.makedirs(args.out, exist_ok=True)
    out = os.path.join(args.out, "bloomie_behaviour_scorecard.xlsx")
    wb.save(out)
    print(f"PRA: {lean(P)} | PASCA: {lean(Q)}  ({shift})")
    print("✓ Excel:", out)


if __name__ == "__main__":
    main()
