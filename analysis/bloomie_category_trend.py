"""
BLOOMIE — Tren revenue bulanan per kategori produk (one-off, TERPISAH dari dashboard).
Juli dipecah: Jul(1-18) pra-reopening vs Jul(19-31) pasca-reopening (reopening 19/07/2026).

Definisi sama dashboard: Revenue = New_Row_Total, kategori = Kategori, channel BLOOMIE
via branch_group. Data dibaca dari SQLite lokal (data_local/transactions.db).

Output Excel ke folder DI LUAR repo (default ~/Bloomie_Analysis/).
  python analysis/bloomie_category_trend.py
  python analysis/bloomie_category_trend.py --reopen 2026-07-19 --topn 10 --out "D:/Analisis"
"""
import os, sys, argparse
from collections import defaultdict
from datetime import date, datetime

_HERE = os.path.dirname(os.path.abspath(__file__))
_BACKEND = os.path.join(os.path.dirname(_HERE), "backend")
sys.path.insert(0, _BACKEND)
try:
    from dotenv import load_dotenv
    load_dotenv(os.path.join(_BACKEND, ".env"))
except Exception:
    pass

import config
import local_store
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

REV = config.REVENUE_COLUMN
MONTHS = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]


def _d(v):
    try:
        return datetime.fromisoformat(str(v)[:10]).date()
    except Exception:
        return None


def period_key(d, reopen):
    """Bucket periode: bulan kalender, kecuali bulan reopening dipecah pra/pasca."""
    m = d.month
    if m == reopen.month and d.year == reopen.year:
        return (m, 0) if d.day < reopen.day else (m, 1)
    return (m, -1)   # -1 = bulan utuh


def period_label(key, reopen):
    m, half = key
    if half == -1:
        return MONTHS[m - 1]
    if half == 0:
        return f"{MONTHS[m-1]} (1-{reopen.day-1}) · pra"
    return f"{MONTHS[m-1]} ({reopen.day}-akhir) · pasca"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--reopen", default="2026-07-19")
    ap.add_argument("--topn", type=int, default=10)
    ap.add_argument("--out", default=os.path.join(os.path.expanduser("~"), "Bloomie_Analysis"))
    args = ap.parse_args()
    reopen = datetime.fromisoformat(args.reopen).date()

    print("Membaca BLOOMIE dari SQLite lokal...")
    conn = local_store.connect(); local_store.ensure_schema(conn)
    cw = local_store.channel_where("BLOOMIE")
    rows = [r for r in local_store.iter_rows(conn, where=cw[0], params=cw[1])
            if config.branch_group(r.get("branch")) == "BLOOMIE"]
    conn.close()
    print(f"  {len(rows)} baris BLOOMIE.")
    if not rows:
        print("Tidak ada data BLOOMIE. Jalankan local_ingest.py dulu."); return

    # kategori top-N by revenue
    kat_rev = defaultdict(float)
    for r in rows:
        kat_rev[r.get("kategori") or "—"] += float(r.get(REV) or 0)
    top = [k for k, _ in sorted(kat_rev.items(), key=lambda x: -x[1])[:args.topn]]
    cols = top + ["Lainnya"]

    def catcol(k):
        return k if k in top else "Lainnya"

    # matriks periode x kategori
    cell = defaultdict(lambda: defaultdict(float))   # periodkey -> cat -> rev
    pra = defaultdict(float); pasca = defaultdict(float)
    for r in rows:
        d = _d(r.get("posting_date"))
        if not d:
            continue
        rev = float(r.get(REV) or 0); c = catcol(r.get("kategori") or "—")
        cell[period_key(d, reopen)][c] += rev
        (pra if d < reopen else pasca)[c] += rev

    def order(k):
        m, half = k
        return (m, 0 if half == -1 else (1 if half == 0 else 2))
    periods = sorted(cell.keys(), key=order)

    # ---------- tulis Excel ----------
    HDR = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="D31137")
    REOP = PatternFill("solid", fgColor="FFF3B0"); BOLD = Font(bold=True)
    NUM = '#,##0'

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Trend per Kategori"
    ws.append(["BLOOMIE — Revenue bulanan per kategori (reopening " + reopen.isoformat() + ")"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    head = ["Periode"] + cols + ["TOTAL"]
    ws.append(head)
    for c in ws[ws.max_row]:
        c.font = HDR; c.fill = FILL

    def write_row(label, data_by_cat, bold=False):
        tot = sum(data_by_cat.get(c, 0) for c in cols)
        ws.append([label] + [round(data_by_cat.get(c, 0)) for c in cols] + [round(tot)])
        for i, c in enumerate(ws[ws.max_row]):
            if i > 0:
                c.number_format = NUM
            if bold:
                c.font = BOLD

    prev_half = None
    for k in periods:
        # sisipkan garis REOPENING antara pra (half 0) dan pasca (half 1)
        if k[1] == 1 and prev_half == 0:
            ws.append([f"———  REOPENING {reopen.isoformat()}  ———"])
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=len(head))
            cc = ws.cell(ws.max_row, 1); cc.fill = REOP; cc.font = BOLD
            cc.alignment = Alignment(horizontal="center")
        write_row(period_label(k, reopen), cell[k])
        prev_half = k[1]

    # total kolom
    totals = defaultdict(float)
    for k in periods:
        for c in cols:
            totals[c] += cell[k].get(c, 0)
    write_row("TOTAL", totals, bold=True)

    # ---------- sheet 2: ringkasan pra vs pasca per kategori ----------
    ws2 = wb.create_sheet("Pra vs Pasca")
    ws2.append(["Ringkasan agregat per kategori: sebelum vs sesudah reopening"]); ws2["A1"].font = Font(bold=True, size=12)
    ws2.append([f"Pra = sebelum {reopen.isoformat()} | Pasca = {reopen.isoformat()} dst (seluruh bulan yang ada)"])
    ws2.append([])
    ws2.append(["Kategori", "Revenue PRA", "Revenue PASCA", "Perubahan %"])
    for c in ws2[ws2.max_row]:
        c.font = HDR; c.fill = FILL
    allcats = sorted(set(list(pra.keys()) + list(pasca.keys())),
                     key=lambda c: -(pra.get(c, 0) + pasca.get(c, 0)))
    for c in allcats:
        a, b = pra.get(c, 0), pasca.get(c, 0)
        chg = f"{(b-a)/a*100:+.1f}%" if a else "—"
        ws2.append([c, round(a), round(b), chg])
        ws2.cell(ws2.max_row, 2).number_format = NUM; ws2.cell(ws2.max_row, 3).number_format = NUM
    ta, tb = sum(pra.values()), sum(pasca.values())
    ws2.append(["TOTAL", round(ta), round(tb), f"{(tb-ta)/ta*100:+.1f}%" if ta else "—"])
    for c in ws2[ws2.max_row]:
        c.font = BOLD

    # lebar kolom
    for w in wb.worksheets:
        for col in w.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            try:
                w.column_dimensions[col[0].column_letter].width = min(26, max(11, width + 2))
            except Exception:
                pass

    os.makedirs(args.out, exist_ok=True)
    out_file = os.path.join(args.out, "bloomie_trend_kategori_bulanan.xlsx")
    wb.save(out_file)
    print(f"\nPeriode: {', '.join(period_label(k, reopen) for k in periods)}")
    print(f"✓ Excel tersimpan di: {out_file}")


if __name__ == "__main__":
    main()
