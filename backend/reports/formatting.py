"""
Helper format standar untuk semua report (aturan #5 & #6 PROMPT_CLAUDE_CODE_REPORT_CENTER).
Rupiah titik ribuan, persen 1 desimal, warna kolom % (merah/kuning/hijau), header standar,
freeze panes. Semua teks Bahasa Indonesia.
"""
from datetime import date
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# --- palet ---
HDR_FILL = PatternFill("solid", fgColor="C00000")          # satu warna header (merah)
HDR_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=13)
SUB_FONT = Font(italic=True, size=9, color="666666")
BOLD = Font(bold=True)
WARN_FILL = PatternFill("solid", fgColor="FFF3B0")         # peringatan cakupan (kuning muda)
GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="DDDDDD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FMT_RUPIAH = '#,##0'
FMT_PERSEN = '0.0%'          # nilai disimpan sbg pecahan (0.123), tampil 12.3%


def pct_fill(value, good=0.9, warn=0.7):
    """Warna sel % pencapaian: >=good hijau, >=warn kuning, sisanya merah. value = pecahan."""
    if value is None or value == "":
        return None
    try:
        v = float(value)
    except Exception:
        return None
    return GREEN if v >= good else (YELLOW if v >= warn else RED)


def write_standard_header(ws, title, periode_dari, periode_sampai, filters: dict,
                          coverage_warnings=None, start_row=1):
    """Tulis header standar (judul, periode, filter aktif, tanggal dibuat, peringatan cakupan).
    Return baris berikutnya yang kosong (untuk mulai menaruh tabel)."""
    r = start_row
    ws.cell(r, 1, title).font = TITLE_FONT; r += 1
    ws.cell(r, 1, f"Periode: {periode_dari} s/d {periode_sampai}").font = SUB_FONT; r += 1
    if filters:
        ftxt = " | ".join(f"{k}: {v}" for k, v in filters.items())
        ws.cell(r, 1, f"Filter aktif: {ftxt}").font = SUB_FONT; r += 1
    ws.cell(r, 1, f"Dibuat: {date.today().isoformat()}").font = SUB_FONT; r += 1
    for w in (coverage_warnings or []):
        c = ws.cell(r, 1, f"⚠ {w}"); c.font = Font(bold=True, color="9C5700"); c.fill = WARN_FILL
        r += 1
    return r + 1   # satu baris kosong pemisah


def style_table_header(ws, row, ncol):
    for j in range(1, ncol + 1):
        c = ws.cell(row, j); c.font = HDR_FONT; c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center"); c.border = BORDER


def rupiah(cell):
    cell.number_format = FMT_RUPIAH


def persen(cell):
    cell.number_format = FMT_PERSEN
