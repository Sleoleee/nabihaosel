"""
Report Center — File 2: Customer Performance.
Sheet 'Ringkasan Customer' (baris induk tebal + anak bisa dilipat; grand total hanya
anak + individu) + 'Detail Customer x Kategori' (lebar, AutoFilter).
Sumber: agg_customer_month, agg_customer_category_month, dim_customer, customer_group.
Status & tier induk dihitung dari agregat (aturan emas), bukan di frontend.
"""
import io
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import settings_store
from utils.calculations import get_customer_tier
from reports import formatting as F
from reports.salesperson import _period_months, _mlabel, MONTH_ABBR


def _status(days):
    if days is None:
        return "AKTIF"
    if days <= 90:
        return "AKTIF"
    if days <= 180:
        return "TIDAK AKTIF"
    return "HILANG"


def build(db, dari, sampai, channels=None, coverage_warnings=None):
    months = _period_months(dari, sampai)
    monthset = set(months)
    years = sorted({y for y, _ in months})
    chan_set = set(channels) if channels else None

    acm = db.table("agg_customer_month").select(
        "customer_code,tahun,bulan,channel,revenue,bills").limit(1000000).execute().data or []
    dimc = {r["customer_code"]: r for r in (db.table("dim_customer").select(
        "customer_code,customer_name,tier,avg_spending_per_month_active,total_revenue,"
        "last_order_date,days_since_last_order,jumlah_bills").limit(1000000).execute().data or [])}
    group_map = settings_store.get_customer_group_map()   # {code: group induk}

    # period revenue/bills per customer + lifetime active-months per customer (utk induk tier)
    per_rev = defaultdict(float); per_bills = defaultdict(int)
    life_rev = defaultdict(float); life_months = defaultdict(set)
    for r in acm:
        c = r["customer_code"]; y, m = int(r["tahun"]), int(r["bulan"])
        rev = float(r.get("revenue") or 0)
        life_rev[c] += rev; life_months[c].add((y, m))
        if (y, m) in monthset and (not chan_set or r.get("channel") in chan_set):
            per_rev[c] += rev; per_bills[c] += int(r.get("bills") or 0)

    active = [c for c in per_rev if per_rev[c] > 0]

    # kelompokkan: induk -> anak; individu = tanpa induk
    groups = defaultdict(list)
    individuals = []
    for c in active:
        g = group_map.get(c)
        (groups[g].append(c) if g else individuals.append(c))

    def cust_row(c):
        d = dimc.get(c, {})
        days = d.get("days_since_last_order")
        return {
            "kode": c, "nama": d.get("customer_name") or c, "induk": group_map.get(c) or "",
            "status": _status(days), "tier": d.get("tier") or "",
            "revenue": round(per_rev[c]),
            "avg_spm": round(float(d.get("avg_spending_per_month_active") or 0)),
            "bills": int(d.get("jumlah_bills") or per_bills[c]),
            "terakhir": d.get("last_order_date") or "", "hari_sejak": days if days is not None else "",
        }

    def induk_row(g, members):
        lifetime_rev = sum(life_rev[c] for c in members)
        amonths = set().union(*[life_months[c] for c in members]) if members else set()
        avg = lifetime_rev / len(amonths) if amonths else 0
        days_list = [dimc.get(c, {}).get("days_since_last_order") for c in members]
        days_list = [x for x in days_list if x is not None]
        dmin = min(days_list) if days_list else None
        last = max((dimc.get(c, {}).get("last_order_date") or "") for c in members)
        return {
            "kode": "", "nama": f"▸ {g}", "induk": g, "status": _status(dmin),
            "tier": get_customer_tier(avg), "revenue": round(sum(per_rev[c] for c in members)),
            "avg_spm": round(avg), "bills": sum(int(dimc.get(c, {}).get("jumlah_bills") or 0) for c in members),
            "terakhir": last, "hari_sejak": dmin if dmin is not None else "",
        }

    wb = Workbook()
    ws = wb.active; ws.title = "Ringkasan Customer"
    filt = {"Channel": ", ".join(channels) if channels else "Semua"}
    start = F.write_standard_header(ws, "Customer Performance — Ringkasan", dari, sampai, filt, coverage_warnings)
    head = ["Kode", "Nama Customer", "Grup Induk", "Status", "Tier", "Total Revenue",
            "Avg Spending/Bulan Aktif", "Bills", "Terakhir Beli", "Hari Sejak Terakhir Beli"]
    hrow = start
    for j, h in enumerate(head, 1):
        ws.cell(hrow, j, h)
    F.style_table_header(ws, hrow, len(head))
    ws.sheet_properties.outlinePr.summaryBelow = False

    def write_cust(r, d, bold=False, level=0):
        ws.cell(r, 1, d["kode"]); ws.cell(r, 2, d["nama"]); ws.cell(r, 3, d["induk"])
        ws.cell(r, 4, d["status"]); ws.cell(r, 5, d["tier"])
        F.rupiah(ws.cell(r, 6, d["revenue"])); F.rupiah(ws.cell(r, 7, d["avg_spm"]))
        ws.cell(r, 8, d["bills"]); ws.cell(r, 9, d["terakhir"]); ws.cell(r, 10, d["hari_sejak"])
        if bold:
            for c in ws[r]: c.font = F.BOLD
        if level:
            ws.row_dimensions[r].outlineLevel = level
            ws.row_dimensions[r].hidden = False

    r = hrow + 1
    gt_rev = 0; gt_bills = 0
    # induk (urut revenue periode desc)
    order = sorted(groups.items(), key=lambda kv: -sum(per_rev[c] for c in kv[1]))
    for g, members in order:
        write_cust(r, induk_row(g, members), bold=True); r += 1
        for c in sorted(members, key=lambda x: -per_rev[x]):
            d = cust_row(c); write_cust(r, d, level=1); r += 1
            gt_rev += d["revenue"]; gt_bills += d["bills"]
    # individu
    for c in sorted(individuals, key=lambda x: -per_rev[x]):
        d = cust_row(c); write_cust(r, d); r += 1
        gt_rev += d["revenue"]; gt_bills += d["bills"]
    # grand total (anak + individu saja)
    ws.cell(r, 2, "GRAND TOTAL (anak + individu)").font = F.BOLD
    F.rupiah(ws.cell(r, 6, gt_rev)); ws.cell(r, 8, gt_bills)
    for c in ws[r]: c.font = F.BOLD
    ws.auto_filter.ref = f"A{hrow}:{get_column_letter(len(head))}{r-1}"
    ws.freeze_panes = f"A{hrow+1}"
    for col, w in zip("ABCDEFGHIJ", [12, 28, 18, 12, 8, 16, 18, 9, 13, 14]):
        ws.column_dimensions[col].width = w

    # ===== Detail Customer x Kategori x Bulan =====
    acckm = db.table("agg_customer_category_month").select(
        "customer_code,kategori,tahun,bulan,channel,revenue").in_("tahun", years).limit(1000000).execute().data or []
    detail = defaultdict(lambda: defaultdict(float))   # (induk,cust,kat) -> ym -> rev
    for r2 in acckm:
        c = r2["customer_code"]; y, m = int(r2["tahun"]), int(r2["bulan"])
        if (y, m) not in monthset or (chan_set and r2.get("channel") not in chan_set):
            continue
        if per_rev.get(c, 0) <= 0:
            continue
        g = group_map.get(c) or "(tanpa induk)"
        detail[(g, c, r2["kategori"])][(y, m)] += float(r2.get("revenue") or 0)

    ws2 = wb.create_sheet("Detail Customer x Kategori")
    start2 = F.write_standard_header(ws2, "Detail Customer × Kategori × Bulan", dari, sampai, filt, coverage_warnings)
    head2 = ["Grup Induk", "Customer", "Kategori"] + [_mlabel(y, m) for y, m in months] + ["Total"]
    hrow2 = start2
    for j, h in enumerate(head2, 1):
        ws2.cell(hrow2, j, h)
    F.style_table_header(ws2, hrow2, len(head2))
    rr = hrow2 + 1
    for (g, c, kat), mm in sorted(detail.items(), key=lambda kv: (kv[0][0], kv[0][1], -sum(kv[1].values()))):
        nm = dimc.get(c, {}).get("customer_name") or c
        ws2.cell(rr, 1, g); ws2.cell(rr, 2, f"{c} — {nm}"); ws2.cell(rr, 3, kat)
        for j, (y, m) in enumerate(months, 4):
            F.rupiah(ws2.cell(rr, j, round(mm.get((y, m), 0))))
        F.rupiah(ws2.cell(rr, 3 + len(months) + 1, round(sum(mm.values()))))
        rr += 1
    ws2.auto_filter.ref = f"A{hrow2}:{get_column_letter(len(head2))}{rr-1}"
    ws2.freeze_panes = f"D{hrow2+1}"
    for col, w in zip("ABC", [18, 28, 20]):
        ws2.column_dimensions[col].width = w

    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    return bio.getvalue()
