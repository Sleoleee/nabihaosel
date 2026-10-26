"""
Report Center — File 4: KPI Manager.
Isi kolom Realisasi di template KPI asli (Review - KPI 2 Managers.xlsx) untuk KPI yang
bisa dihitung dari data penjualan; pertahankan semua rumus & format. Sel otomatis diberi
warna hijau + comment; bila data channel belum masuk, warna merah + catatan.
Tambah sheet 'Rincian Perhitungan'. Manager: 'branding' / 'commerce' / 'both'.
"""
import io, os, json
from datetime import date
from collections import defaultdict
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.comments import Comment

import config

_TPL_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "report_templates")
TEMPLATE = os.path.join(_TPL_DIR, "Review - KPI 2 Managers.xlsx")
CONFIG = os.path.join(_TPL_DIR, "kpi_config.json")

GREEN = PatternFill("solid", fgColor="C6EFCE")
PINK = PatternFill("solid", fgColor="FFC7CE")
BULAN_ID = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
            "Agustus", "September", "Oktober", "November", "Desember"]
BRAND_CAT = {"Kamino": "AKSESORIS RAMBUT KAMINO", "Roche": "ROCHE", "Glimore": "GLIMORE",
             "Mimori": "MIMORI", "Others": "BLISS BEAUTY"}
SHEET_BY_MANAGER = {"branding": "Branding & Digital Marketing", "commerce": "Commerce & Offline Store"}


def _prev(ym):
    y, m = ym; m -= 1
    if m == 0: m = 12; y -= 1
    return (y, m)


def _prior3(ym):
    out = set(); y, m = ym
    for _ in range(3):
        m -= 1
        if m == 0: m = 12; y -= 1
        out.add((y, m))
    return out


def compute_metrics(db, ym):
    """Hitung semua metrik KPI otomatis untuk bulan penilaian ym=(tahun,bulan)."""
    y, mo = ym
    out = {}; detail = {}

    # --- omset TOKO/BRAND/LAINNYA dari agg_revenue_bucket_month ---
    rb = db.table("agg_revenue_bucket_month").select(
        "kelompok,tahun,bulan,channel,revenue").eq("tahun", y).eq("bulan", mo).limit(1000000).execute().data or []
    bucket = defaultdict(float); bucket_ch = defaultdict(lambda: defaultdict(float))
    for r in rb:
        bucket[r["kelompok"]] += float(r.get("revenue") or 0)
        bucket_ch[r["kelompok"]][r.get("channel")] += float(r.get("revenue") or 0)
    out["omset_perusahaan"] = sum(bucket.values())
    out["omset_brand"] = bucket.get("BRAND", 0.0)
    out["omset_toko"] = bucket.get("TOKO", 0.0)
    detail["omset_per_kelompok"] = dict(bucket)
    detail["omset_per_channel"] = {k: dict(v) for k, v in bucket_ch.items()}

    # --- GMV MoM (E-Commerce) dari agg_customer_category_month (kategori = brand) ---
    years = sorted({y, _prev(ym)[0]})
    accm = db.table("agg_customer_category_month").select(
        "kategori,tahun,bulan,channel,revenue").in_("tahun", years).limit(1000000).execute().data or []
    ecom = defaultdict(float)   # (kategori, (y,m)) -> rev di E-Commerce
    for r in accm:
        if r.get("channel") != "E-Commerce":
            continue
        ecom[(str(r["kategori"]).upper(), (int(r["tahun"]), int(r["bulan"])))] += float(r.get("revenue") or 0)
    pm = _prev(ym)
    mim_now = ecom.get((BRAND_CAT["Mimori"], ym), 0.0); mim_prev = ecom.get((BRAND_CAT["Mimori"], pm), 0.0)
    out["gmv_mom_Mimori"] = ((mim_now - mim_prev) / mim_prev) if mim_prev else None
    other_cats = {BRAND_CAT[b] for b in ("Roche", "Glimore", "Kamino", "Others")}
    bl_now = sum(ecom.get((c, ym), 0.0) for c in other_cats); bl_prev = sum(ecom.get((c, pm), 0.0) for c in other_cats)
    out["gmv_mom_BrandLain"] = ((bl_now - bl_prev) / bl_prev) if bl_prev else None
    detail["gmv_ecommerce"] = {"Mimori": {"ini": round(mim_now), "lalu": round(mim_prev)},
                               "BrandLain": {"ini": round(bl_now), "lalu": round(bl_prev)}}

    # --- Brand Reorder Rate def(b): pembeli 3 bln sebelum M yang beli lagi di M ---
    P3 = _prior3(ym)
    byears = sorted({yy for (yy, _) in (P3 | {ym})})
    bc = db.table("agg_customer_category_month").select(
        "customer_code,kategori,tahun,bulan").in_("tahun", byears).limit(1000000).execute().data or []
    brandcust = defaultdict(lambda: defaultdict(set))   # kategori_brand -> ym -> set(cust)
    catset = {v for v in BRAND_CAT.values()}
    for r in bc:
        k = str(r["kategori"]).upper()
        if k in catset:
            brandcust[k][(int(r["tahun"]), int(r["bulan"]))].add(r["customer_code"])
    reorder_detail = {}
    for brand, cat in BRAND_CAT.items():
        prior_buyers = set().union(*[brandcust[cat].get(p, set()) for p in P3]) if P3 else set()
        now_buyers = brandcust[cat].get(ym, set())
        rate = (len(prior_buyers & now_buyers) / len(prior_buyers)) if prior_buyers else None
        out[f"reorder_{brand}"] = rate
        reorder_detail[brand] = {"pembeli_3bln": len(prior_buyers), "reorder": len(prior_buyers & now_buyers)}
    detail["reorder"] = reorder_detail
    return out, detail


def build(db, bulan, manager="both"):
    y, mo = int(bulan[:4]), int(bulan[5:7]); ym = (y, mo)
    cfg = json.load(open(CONFIG, encoding="utf-8"))
    wb = openpyxl.load_workbook(TEMPLATE)
    metrics, detail = compute_metrics(db, ym)
    today = date.today().isoformat()

    want = list(SHEET_BY_MANAGER.values()) if manager == "both" else [SHEET_BY_MANAGER[manager]]
    for sn in list(wb.sheetnames):
        if sn in cfg["sheets"] and sn not in want:
            wb.remove(wb[sn])

    omset_kosong = metrics.get("omset_perusahaan", 0) <= 0
    for sn in want:
        sc = cfg["sheets"][sn]; ws = wb[sn]
        ws[sc["period_month_cell"]] = BULAN_ID[mo - 1]
        ws[sc["period_year_cell"]] = y
        col = sc["realisasi_col"]; ccol = sc["catatan_col"]
        for kno, kd in sc["kpi"].items():
            val = metrics.get(kd["metric"])
            cell = ws[f"{col}{kd['row']}"]
            if val is None:
                cell.value = "n/a"; cell.fill = PINK
                ws[f"{ccol}{kd['row']}"] = "Data belum cukup untuk dihitung otomatis."
                continue
            cell.value = round(val) if kd["metric"].startswith("omset") else round(val, 4)
            cell.fill = PINK if omset_kosong and kd["metric"].startswith("omset") else GREEN
            note = f"Otomatis dari dashboard, {today}."
            if kd.get("proxy"):
                note += " Proxy dari data SAP, bukan GMV Seller Center."
            if omset_kosong and kd["metric"].startswith("omset"):
                note += " PERINGATAN: sebagian data channel bulan ini mungkin belum masuk."
            try:
                cell.comment = Comment(note, "Dashboard")
            except Exception:
                pass

    # --- sheet Rincian Perhitungan ---
    rc = wb.create_sheet("Rincian Perhitungan")
    rc.append(["RINCIAN PERHITUNGAN KPI OTOMATIS", f"Bulan: {BULAN_ID[mo-1]} {y}", f"Dibuat: {today}"])
    rc.append([])
    rc.append(["Omset per kelompok"])
    for k, v in detail["omset_per_kelompok"].items():
        rc.append(["", k, round(v)])
    rc.append(["Omset per channel"])
    for kel, chv in detail["omset_per_channel"].items():
        for ch, v in chv.items():
            rc.append(["", f"{kel} / {ch}", round(v)])
    rc.append([])
    rc.append(["GMV E-Commerce (bulan ini vs lalu)"])
    for b, vv in detail["gmv_ecommerce"].items():
        rc.append(["", b, f"ini {vv['ini']:,}", f"lalu {vv['lalu']:,}"])
    rc.append([])
    rc.append(["Brand Reorder Rate def(b): pembeli 3 bln sebelum -> reorder bulan ini"])
    for b, vv in detail["reorder"].items():
        rc.append(["", b, f"pembeli_3bln {vv['pembeli_3bln']}", f"reorder {vv['reorder']}"])
    rc.append([])
    rc.append(["Catatan: KPI 6.1/7.1 (SPV/Team Achievement) perlu pemetaan manager->grup SPV; "
               "diisi manual sampai pemetaan ditetapkan."])
    for c in rc.columns:
        try: rc.column_dimensions[c[0].column_letter].width = 26
        except Exception: pass

    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    return bio.getvalue()
