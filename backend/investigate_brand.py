"""
Pelacak sumber BRAND — cari di mana brand sebenarnya tercatat.
Read-only. Tiga sudut: (A) header kolom file sumber, (B) Kategori & prefix kode
item dari SQLite, (C) sidik jari item brand yang sudah ketahuan (via keyword).

  python investigate_brand.py ["file2025.xlsx" "file2026.xlsx" ...]

File opsional hanya untuk membaca HEADER kolom (cek apakah ada kolom brand/U_Kategori
/Item Group yang belum di-ingest). Analisis SQLite selalu jalan.
"""
import sys, os, re
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config
import local_store
import openpyxl

REV = config.REVENUE_COLUMN
BRAND_KEYWORDS = {"Mimori": ["MIMORI"], "Roche": ["ROCHE"], "Glimore": ["GLIMORE", "GLIMOR"],
                  "Bliss Beauty": ["BLISS"], "Kamino": ["KAMINO"]}


def brand_of(desc, item_no):
    s = (str(desc or "") + " " + str(item_no or "")).upper()
    for b, kws in BRAND_KEYWORDS.items():
        if any(k in s for k in kws):
            return b
    return None


def code_prefix(item_no):
    s = str(item_no or "").strip().upper()
    m = re.match(r"^[A-Z]+", s)          # huruf awal (mis. KKJ, TKG, LK)
    return m.group(0) if m else (s.split("-")[0] if s else "?")


def qf(v):
    try: return float(v)
    except Exception: return 0.0


def dump_headers(path):
    print("\n" + "-" * 68); print(f"HEADER: {os.path.basename(path)}")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        print(f"  ! gagal: {str(e)[:60]}"); return
    for sn in wb.sheetnames:
        ws = wb[sn]
        try:
            hdr = next(ws.iter_rows(values_only=True))
        except StopIteration:
            continue
        cols = [str(h).strip() for h in hdr if h is not None]
        # tandai kolom yang mungkin brand
        hits = [c for c in cols if re.search(r"brand|u_?kategori|item.?group|grup|merk|principal", c, re.I)]
        print(f"  sheet '{sn}' ({len(cols)} kolom): {cols}")
        if hits:
            print(f"    >>> KANDIDAT BRAND: {hits}")
    wb.close()


def main():
    files = sys.argv[1:]
    for f in files:
        if os.path.exists(f):
            dump_headers(f)
        else:
            print(f"  ! lewati (tak ada): {f}")

    conn = local_store.connect(); local_store.ensure_schema(conn)
    kat_rev = defaultdict(float)
    pref_rev = defaultdict(float)
    brand_kat = defaultdict(lambda: defaultdict(float))     # brand -> kategori -> rev
    brand_pref = defaultdict(lambda: defaultdict(float))    # brand -> prefix -> rev
    item_seen = {}
    total = 0.0
    for r in local_store.iter_rows(conn, cols=["item_no", "item_description", "kategori", REV],
                                   where="CAST(year AS INTEGER) >= ?", params=(config.MIN_YEAR,)):
        rev = qf(r.get(REV)); total += rev
        kat = (r.get("kategori") or "—")
        it = r.get("item_no")
        kat_rev[kat] += rev
        pref_rev[code_prefix(it)] += rev
        if it not in item_seen:
            item_seen[it] = brand_of(r.get("item_description"), it)
        b = item_seen[it]
        if b:
            brand_kat[b][kat] += rev
            brand_pref[b][code_prefix(it)] += rev
    conn.close()
    total = total or 1

    print("\n" + "=" * 68); print("(B1) KATEGORI teratas by revenue (top 40) — apakah brand = kategori?"); print("=" * 68)
    for kat, rv in sorted(kat_rev.items(), key=lambda x: -x[1])[:40]:
        print(f"  {rv/total*100:5.2f}%  Rp{round(rv):>16,}  {kat}")

    print("\n" + "=" * 68); print("(B2) PREFIX kode item teratas (top 40) — apakah brand = prefix kode?"); print("=" * 68)
    for p, rv in sorted(pref_rev.items(), key=lambda x: -x[1])[:40]:
        print(f"  {rv/total*100:5.2f}%  Rp{round(rv):>16,}  {p}")

    print("\n" + "=" * 68); print("(C) SIDIK JARI item brand yang ketahuan (via keyword) — Kategori & prefix dominan"); print("=" * 68)
    for b in BRAND_KEYWORDS:
        kats = sorted(brand_kat[b].items(), key=lambda x: -x[1])[:5]
        prefs = sorted(brand_pref[b].items(), key=lambda x: -x[1])[:5]
        print(f"\n  {b}:")
        print(f"    Kategori dominan: {[(k, round(v)) for k, v in kats]}")
        print(f"    Prefix dominan  : {[(p, round(v)) for p, v in prefs]}")
    print("\nSelesai (read-only).")


if __name__ == "__main__":
    main()
