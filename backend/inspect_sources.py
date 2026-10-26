"""
Inspektor file sumber — lihat isi file Excel transaksi SEBELUM di-ingest,
supaya tahu channel/tahun/tanggal yang terkandung & menghindari dobel.

  python inspect_sources.py "file1.xlsx" "file2.xlsx" ...
  python inspect_sources.py --sheet Bloomie "file.xlsx"   # batasi ke sheet tertentu

Tidak menyentuh database apa pun. Read-only.
"""
import sys, os, io
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config
import openpyxl
from utils.parser import parse_excel


def list_sheets(path):
    """Daftar nama sheet + apakah sheet transaksi (punya Posting Date + New_Row_Total)."""
    from utils.parser import COLUMN_MAP
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        print(f"  ! gagal buka: {str(e)[:70]}"); return
    print("  SHEET:")
    for sn in wb.sheetnames:
        ws = wb[sn]
        try:
            header = next(ws.iter_rows(values_only=True))
        except StopIteration:
            print(f"    - {sn}  (kosong)"); continue
        cols = {COLUMN_MAP.get(str(h).strip()) for h in header if h is not None}
        is_tx = ("posting_date" in cols) and ("new_row_total" in cols)
        print(f"    - {sn}  {'[TRANSAKSI]' if is_tx else '(bukan transaksi / dilewati)'}")
    wb.close()


def summarize(path, only_sheet=None, exclude_sheet=None):
    print("\n" + "=" * 70)
    print(f"FILE: {os.path.basename(path)}"
          + (f"   [sheet~'{only_sheet}']" if only_sheet else "")
          + (f"   [exclude~'{exclude_sheet}']" if exclude_sheet else ""))
    print("=" * 70)
    list_sheets(path)
    try:
        with open(path, "rb") as f:
            by_year = parse_excel(io.BytesIO(f.read()), only_sheet=only_sheet, exclude_sheet=exclude_sheet)
    except FileNotFoundError:
        print("  ! FILE TIDAK DITEMUKAN"); return
    except Exception as e:
        print(f"  ! gagal parse: {str(e)[:80]}"); return
    if not by_year:
        print("  (tak ada sheet transaksi valid)"); return

    for year in sorted(by_year):
        recs, skipped = by_year[year]
        dates = [r.get("posting_date") for r in recs if r.get("posting_date")]
        dmin = min(dates) if dates else "-"
        dmax = max(dates) if dates else "-"
        # distribusi channel
        ch_nota = defaultdict(set); ch_rows = defaultdict(int)
        for r in recs:
            ch = config.branch_group(r.get("branch"))
            ch_rows[ch] += 1
            if r.get("document_number") is not None:
                ch_nota[ch].add(r.get("document_number"))
        print(f"\n  TAHUN {year}: {len(recs):,} baris | tanggal {dmin} s/d {dmax} | (skip tahun lain: {skipped:,})")
        print(f"    {'Channel':<16}{'Baris':>10}{'Nota unik':>12}")
        print(f"    {'-'*38}")
        for ch in config.BRANCH_GROUP_ORDER:
            if ch_rows.get(ch):
                print(f"    {ch:<16}{ch_rows[ch]:>10,}{len(ch_nota[ch]):>12,}")


def main():
    args = sys.argv[1:]
    only_sheet = exclude_sheet = None
    if "--sheet" in args:
        i = args.index("--sheet"); only_sheet = args[i + 1]; del args[i:i + 2]
    if "--exclude-sheet" in args:
        i = args.index("--exclude-sheet"); exclude_sheet = args[i + 1]; del args[i:i + 2]
    if not args:
        print("Pakai: python inspect_sources.py [--sheet NAMA] [--exclude-sheet NAMA] \"file1.xlsx\" ..."); return
    for path in args:
        summarize(path, only_sheet, exclude_sheet)
    print("\n" + "=" * 70)
    print("Selesai (read-only, DB tidak disentuh).")


if __name__ == "__main__":
    main()
